"""Saved S3 knowledge cards flip through internal PPTX slide links."""

from io import BytesIO

from pptx import Presentation
from pptx.util import Inches

from ai_engine.gen.export_ppt import ExportPPTConfig, PPTExporter
from ai_engine.gen.export_quality import validate_pptx


def _slide_text(slide):
    return "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)


def _linked_shape(slide, text):
    return next(shape for shape in slide.shapes if shape.has_text_frame and text in shape.text)


def _card_background(slide):
    return next(shape for shape in slide.shapes
                if shape.width == Inches(11.13) and shape.height == Inches(4.65))


def test_saved_flashcards_have_separate_linked_front_and_back_slides():
    node = {
        "node_id": "card_1",
        "title": "排序算法",
        "type": "quiz",  # stale type must not turn the card into a quiz page
        "interactive_plan": {"kind": "flashcard"},
        "metadata": {"interactive_asset": {"items": [
            {"type": "quiz", "questions": [{"question": "旧答题不应出现"}]},
            {"type": "flashcard", "front": "快速排序", "back": "采用分治思想"},
            {"type": "flashcard", "front": "冒泡排序", "back": "比较相邻元素"},
        ]}},
    }
    exporter = PPTExporter(ExportPPTConfig(
        include_auto_cover=False, include_auto_toc=False,
        include_auto_summary=False, include_closing_slide=False,
    ))
    data = exporter.export_to_bytes([node], "排序算法").getvalue()
    deck = Presentation(BytesIO(data))
    assert len(deck.slides) == 4
    front_1, back_1, front_2, back_2 = deck.slides
    assert "快速排序" in _slide_text(front_1)
    assert "采用分治思想" not in _slide_text(front_1)
    assert "采用分治思想" in _slide_text(back_1)
    assert "快速排序" not in _slide_text(back_1)
    assert "冒泡排序" in _slide_text(front_2)
    assert "比较相邻元素" in _slide_text(back_2)
    assert "旧答题不应出现" not in "\n".join(_slide_text(s) for s in deck.slides)

    # Inspect the serialized package, not just the in-memory Presentation.
    assert _linked_shape(front_1, "快速排序").click_action.target_slide == back_1
    assert _linked_shape(back_1, "采用分治思想").click_action.target_slide == front_1
    assert _card_background(front_1).click_action.target_slide == back_1
    assert _card_background(back_1).click_action.target_slide == front_1
    assert _linked_shape(front_1, "点击卡片查看释义").click_action.target_slide == back_1
    assert _linked_shape(back_2, "点击卡片返回正面").click_action.target_slide == front_2
    assert _linked_shape(back_1, "下一张").click_action.target_slide == front_2
    assert _linked_shape(front_2, "上一张").click_action.target_slide == front_1
    assert validate_pptx(data)["status"] != "failed"
