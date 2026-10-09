"""S4 DOCX export must produce the structured A12 lesson plan."""

import time
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
import pytest

import ai_engine.main as main_module
from ai_engine.gen.lesson_plan import LessonPlanGenerator
from ai_engine.main import app


@pytest.mark.parametrize("source,reason,filled", [
    ("fallback", "AI 教案生成超时", []),
    ("llm", "", ["教学目标"]),
])
def test_docx_export_status_includes_lesson_plan_diagnostics(
    tmp_path, monkeypatch, source, reason, filled,
):
    monkeypatch.chdir(tmp_path)
    plan = LessonPlanGenerator()._build_fallback_plan(
        [{"node_id": "n1", "title": "排序算法", "content": "比较与交换"}],
        {"topic": "排序算法"},
    ).to_dict()
    plan.update(data_source=source, fallback_reason=reason, filled_sections=filled)
    monkeypatch.setattr(main_module, "generate_lesson_plan_cached", lambda *_args: plan)

    with TestClient(app) as client:
        created = client.post("/api/v1/export/lesson-plan-diagnostics", json={
            "format": "docx", "version": 1, "courseware_title": "排序算法",
            "nodes": [{"node_id": "n1", "title": "排序算法", "content": "比较与交换"}],
        })
        assert created.status_code == 200, created.text
        for _ in range(80):
            result = client.get(f"/api/v1/export/status/{created.json()['job_id']}").json()
            if result["status"] in {"done", "failed"}:
                break
            time.sleep(0.025)

    assert result["status"] == "done", result
    assert result["lesson_plan_source"] == source
    assert result["lesson_plan_fallback_reason"] == reason
    assert result["lesson_plan_filled_sections"] == filled


def test_docx_export_contains_generated_sections_and_emphasized_headings(
    tmp_path, monkeypatch
):
    monkeypatch.chdir(tmp_path)

    with TestClient(app) as client:
        created = client.post(
            "/api/v1/export/generation-lesson-plan",
            json={
                "format": "docx",
                "version": 3,
                "courseware_title": "图搜索",
                "nodes": [{
                    "node_id": "selected-v3-node",
                    "title": "广度优先搜索",
                    "type": "content",
                    "content": "使用队列按层访问图中的顶点。",
                    "key_points": ["队列先进先出", "按层遍历"],
                }],
                "generation_instruction_set": {
                    "intent": {
                        "topic": "图搜索",
                        "duration_minutes": 45,
                        "target_student": "大一计算机专业学生",
                        "with_interactive": True,
                    },
                    "confirmed_evidence": {
                        "synthesized_context": "广度优先搜索使用队列实现逐层遍历。"
                    },
                },
            },
        )
        assert created.status_code == 200, created.text
        job_id = created.json()["job_id"]

        for _ in range(80):
            result = client.get(f"/api/v1/export/status/{job_id}").json()
            if result["status"] in {"done", "failed"}:
                break
            time.sleep(0.025)

        assert result["status"] == "done", result
        assert result["version"] == 3
        downloaded = client.get(result["files"]["docx"])
        assert downloaded.status_code == 200

    document = Document(BytesIO(downloaded.content))
    paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
    table_text = "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    text = "\n".join(paragraph.text for paragraph in paragraphs) + "\n" + table_text
    for section in ("教学目标", "教学过程", "教学方法", "课堂活动设计", "课后作业"):
        assert section in text
    assert "大一计算机专业学生" in text

    required_headings = ("三、教学方法", "五、课堂活动设计", "六、课后作业")
    heading_runs = {
        run.text: run
        for paragraph in paragraphs
        for run in paragraph.runs
        if run.text in required_headings
    }
    assert set(heading_runs) == set(required_headings)
    for run in heading_runs.values():
        assert run.bold is True
        assert run.font.size.pt == 14
