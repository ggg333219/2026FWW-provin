"""S4 export quality behavior through the public AI Engine export API."""

import time
import zipfile
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

import ai_engine.main as main_module
import ai_engine.gen.version as version_module
from ai_engine.gen.version import VersionManager
from ai_engine.main import app


class MemoryExportStore:
    def __init__(self):
        self.jobs = {}

    async def save_export(self, job_id, export_data):
        self.jobs[job_id] = export_data

    async def load_export(self, job_id):
        return self.jobs.get(job_id)


def _prepare_generation(tmp_path, monkeypatch, generation_id, snapshots):
    monkeypatch.chdir(tmp_path)

    def versions_dir(gen_id):
        path = tmp_path / "versions" / gen_id
        path.mkdir(parents=True, exist_ok=True)
        return path

    monkeypatch.setattr(version_module, "_versions_dir", versions_dir)
    monkeypatch.setattr(main_module, "get_version_manager", VersionManager)
    store = MemoryExportStore()
    monkeypatch.setattr(main_module, "_task_store", store)
    manager = VersionManager()
    for nodes in snapshots:
        manager.save_snapshot(generation_id, nodes)
    return store


def _poll_export(client, job_id):
    for _ in range(40):
        response = client.get(f"/api/v1/export/status/{job_id}")
        assert response.status_code == 200
        data = response.json()
        if data["status"] in {"done", "failed"}:
            return data
        time.sleep(0.025)
    raise AssertionError("export did not reach a terminal state")


def test_selected_version_export_reports_quality_and_remains_downloadable(
    tmp_path, monkeypatch
):
    store = _prepare_generation(tmp_path, monkeypatch, "generation-quality", [
        [{
            "node_id": "v1-slide",
            "title": "Version One Marker",
            "type": "content",
            "content": "This content belongs to version one.",
        }],
        [{
            "node_id": "v2-slide",
            "title": "Version Two Marker",
            "type": "content",
            "content": "This content belongs to version two.",
        }],
    ])

    with TestClient(app) as client:
        created = client.post(
            "/api/v1/export/generation-quality",
            json={"format": "pptx", "version": 1},
        )
        assert created.status_code == 200
        job_id = created.json()["job_id"]
        result = _poll_export(client, job_id)

        assert result["status"] == "done"
        assert result["quality_report"]["status"] == "passed"
        assert result["quality_report"]["slide_count"] >= 1
        assert result["version"] == 1
        assert result["download_url"]
        downloaded = client.get(result["download_url"])
        assert downloaded.status_code == 200
        with zipfile.ZipFile(BytesIO(downloaded.content)) as archive:
            slide_xml = b" ".join(
                archive.read(name)
                for name in archive.namelist()
                if name.startswith("ppt/slides/slide") and name.endswith(".xml")
            )
        assert b"Version One Marker" in slide_xml
        assert b"Version Two Marker" not in slide_xml
        main_module._export_tasks.pop(job_id)
        restored = client.get(f"/api/v1/export/status/{job_id}").json()
        assert restored["quality_report"] == result["quality_report"]

    assert store.jobs[job_id]["quality_report"]["status"] == "passed"


def _empty_pptx_bytes():
    output = BytesIO()
    Presentation().save(output)
    return output.getvalue()


@pytest.mark.parametrize(
    ("scenario", "finding_code"),
    [("corrupt", "invalid_presentation"), ("empty", "empty_presentation")],
)
def test_invalid_pptx_fails_export_and_cannot_be_downloaded(
    tmp_path, monkeypatch, scenario, finding_code
):
    _prepare_generation(tmp_path, monkeypatch, "generation-invalid", [[{
        "node_id": "invalid-slide", "title": "Invalid", "type": "content",
    }]])
    bad_bytes = b"not a powerpoint file" if scenario == "corrupt" else _empty_pptx_bytes()
    monkeypatch.setattr(
        main_module.PPTExporter,
        "export_to_bytes",
        lambda *_args, **_kwargs: BytesIO(bad_bytes),
    )

    with TestClient(app) as client:
        created = client.post(
            "/api/v1/export/generation-invalid", json={"format": "pptx", "version": 1}
        )
        job_id = created.json()["job_id"]
        result = _poll_export(client, job_id)

        assert result["status"] == "failed"
        assert result["quality_report"]["status"] == "failed"
        assert result["quality_report"]["findings"][0]["code"] == finding_code
        assert result["download_url"] is None
        assert client.get(f"/download/{job_id}.pptx").status_code == 404


def test_text_fit_warning_keeps_valid_pptx_downloadable(tmp_path, monkeypatch):
    _prepare_generation(tmp_path, monkeypatch, "generation-warning", [[{
        "node_id": "warning-slide", "title": "Warning", "type": "content",
    }]])

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(0.2))
    paragraph = text_box.text_frame.paragraphs[0]
    paragraph.text = "This deliberately overlong text cannot fit in this tiny box."
    paragraph.runs[0].font.size = Pt(24)
    text_box.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    pptx_bytes = BytesIO()
    presentation.save(pptx_bytes)
    monkeypatch.setattr(
        main_module.PPTExporter,
        "export_to_bytes",
        lambda *_args, **_kwargs: BytesIO(pptx_bytes.getvalue()),
    )

    with TestClient(app) as client:
        created = client.post(
            "/api/v1/export/generation-warning", json={"format": "pptx", "version": 1}
        )
        job_id = created.json()["job_id"]
        result = _poll_export(client, job_id)

        assert result["status"] == "done"
        assert result["quality_report"]["status"] == "passed_with_warnings"
        assert result["quality_report"]["findings"][0]["code"] == "text_may_overflow"
        assert result["download_url"]
        assert client.get(result["download_url"]).status_code == 200


def test_text_outside_slide_blocks_download(tmp_path, monkeypatch):
    _prepare_generation(tmp_path, monkeypatch, "generation-bounds", [[{
        "node_id": "bounds-slide", "title": "Bounds", "type": "content",
    }]])

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(13.2), Inches(1), Inches(1), Inches(1))
    box.text = "Text beyond the right edge"
    box.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    pptx_bytes = BytesIO()
    presentation.save(pptx_bytes)
    monkeypatch.setattr(
        main_module.PPTExporter,
        "export_to_bytes",
        lambda *_args, **_kwargs: BytesIO(pptx_bytes.getvalue()),
    )

    with TestClient(app) as client:
        created = client.post(
            "/api/v1/export/generation-bounds", json={"format": "pptx", "version": 1}
        )
        job_id = created.json()["job_id"]
        result = _poll_export(client, job_id)

        assert result["status"] == "failed"
        finding = result["quality_report"]["findings"][0]
        assert finding["code"] == "text_outside_slide"
        assert finding["severity"] == "error"
        assert finding["slide_number"] == 1
        assert result["download_url"] is None


def test_validator_failure_fails_closed_with_sanitized_report(tmp_path, monkeypatch):
    _prepare_generation(tmp_path, monkeypatch, "generation-validator-error", [[{
        "node_id": "validator-slide", "title": "Validator", "type": "content",
    }]])

    def unavailable(_data):
        raise RuntimeError("private filesystem path must not escape")

    monkeypatch.setattr(main_module, "validate_pptx", unavailable)

    with TestClient(app) as client:
        created = client.post(
            "/api/v1/export/generation-validator-error",
            json={"format": "pptx", "version": 1},
        )
        job_id = created.json()["job_id"]
        result = _poll_export(client, job_id)

        assert result["status"] == "failed"
        assert result["quality_report"]["findings"][0]["code"] == "validator_unavailable"
        assert "private filesystem path" not in result["error"]
        assert result["download_url"] is None
