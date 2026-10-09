"""S3 已生成翻卡经过版本保存后，S4 HTML5 导出仍须带卡片。"""

import pytest
from io import BytesIO

from pptx import Presentation

from ai_engine.core.schema import TeachingNode
from ai_engine.gen.export_html5 import export_interactive_html, extract_interactive_items
from ai_engine.gen.export_ppt import PPTExporter
from ai_engine.gen.interactive import InteractiveBuilder


def test_export_uses_saved_interactive_asset_items(tmp_path):
    cards = [
        {"type": "flashcard", "front": f"问题 {i}", "back": f"答案 {i}", "tags": ["core"]}
        for i in range(1, 8)
    ]
    node = {
        "node_id": "node_1",
        "title": "定积分的几何意义",
        "interactive_plan": {"kind": "flashcard"},
        "metadata": {"interactive_asset": {"status": "done", "items": cards}},
        "summary": "不应覆盖已经生成的卡片",
    }

    extracted = extract_interactive_items([node])
    assert extracted == cards
    result = export_interactive_html([node], "高等数学", output_dir=str(tmp_path))
    html = (tmp_path / "interactive.html").read_text(encoding="utf-8")
    assert result["flashcard_count"] == 7
    assert result["quiz_count"] == 0
    assert "问题 1" in html and "答案 7" in html
    assert 'var DATA = [];' not in html
    assert "知识翻卡" in html


def test_flashcard_generation_falls_back_when_segments_have_no_knowledge_card():
    items = InteractiveBuilder().build_all({
        "title": "快速排序", "summary": "分治思想",
        "segments": [{"text": "课堂讲解"}],
        "interactive_plan": {"kind": "flashcard"},
    })
    assert items == [{"type": "flashcard", "front": "快速排序", "back": "分治思想", "tags": ["core"]}]


def test_export_keeps_direct_interactive_items():
    node = {
        "title": "排序算法",
        "interactive_items": [{"type": "flashcard", "front": "快速排序", "back": "平均 O(n log n)"}],
    }
    assert extract_interactive_items([node])[0]["back"] == "平均 O(n log n)"


def test_generation_node_model_preserves_interactive_plan_and_items():
    node = TeachingNode.model_validate({
        "node_id": "node_1",
        "title": "快速排序",
        "interactive_plan": {"kind": "flashcard"},
        "interactive_items": [{"type": "flashcard", "front": "分区", "back": "按基准划分数组"}],
    })
    saved = node.model_dump(mode="json")
    assert saved["interactive_plan"]["kind"] == "flashcard"
    assert extract_interactive_items([saved])[0]["back"] == "按基准划分数组"


def test_saved_flashcard_plan_uses_existing_key_points():
    node = {
        "title": "冒泡排序",
        "interactive_plan": {"kind": "flashcard"},
        "key_points": ["比较相邻元素", "每轮将最大元素移到末尾"],
    }
    items = extract_interactive_items([node])
    assert len(items) == 2
    assert [item["back"] for item in items] == node["key_points"]
    assert all(item["type"] == "flashcard" for item in items)


def test_empty_interactive_export_reports_error(tmp_path):
    with pytest.raises(ValueError, match="没有可导出的知识翻卡"):
        export_interactive_html([{"title": "普通讲授页"}], output_dir=str(tmp_path))
    assert not (tmp_path / "interactive.html").exists()


def test_html5_export_excludes_quizzes_even_when_asset_contains_both(tmp_path):
    node = {
        "title": "排序算法",
        "metadata": {"interactive_asset": {"items": [
            {"type": "quiz", "title": "答题挑战", "questions": [{
                "question": "不应出现在翻卡导出中的测验题",
                "options": ["选项 A", "选项 B"], "answer": "A",
            }]},
            {"type": "flashcard", "front": "快速排序", "back": "采用分治思想"},
        ]}},
    }
    result = export_interactive_html([node], output_dir=str(tmp_path))
    html = (tmp_path / "interactive.html").read_text(encoding="utf-8")
    assert result["quiz_count"] == 0
    assert result["flashcard_count"] == 1
    assert "快速排序" in html
    assert "不应出现在翻卡导出中的测验题" not in html
    assert "答题挑战 · 知识翻卡" not in html


def test_quiz_only_asset_cannot_export_as_flashcards(tmp_path):
    node = {"title": "排序算法", "metadata": {"interactive_asset": {"items": [{
        "type": "quiz", "questions": [{
            "question": "哪种排序最快？", "options": ["甲", "乙"], "answer": "A",
        }],
    }]}}}
    with pytest.raises(ValueError, match="没有可导出的知识翻卡"):
        export_interactive_html([node], output_dir=str(tmp_path))


def test_plain_quiz_question_is_not_disguised_as_knowledge_card(tmp_path):
    node = {"title": "排序算法测验", "interactive_questions": ["冒泡排序如何比较元素？"], "summary": "课堂练习"}
    with pytest.raises(ValueError, match="没有可导出的知识翻卡"):
        export_interactive_html([node], output_dir=str(tmp_path))


def test_planned_flashcard_ignores_stale_quiz_and_uses_existing_points():
    node = {
        "title": "冒泡排序",
        "interactive_plan": {"kind": "flashcard"},
        "key_points": ["比较相邻元素"],
        "metadata": {"interactive_asset": {"items": [{
            "type": "quiz", "questions": [{
                "question": "旧测验", "options": ["甲", "乙"], "answer": "A",
            }],
        }]}},
    }
    cards = extract_interactive_items([node], flashcards_only=True)
    assert len(cards) == 1
    assert cards[0]["type"] == "flashcard"
    assert cards[0]["back"] == "比较相邻元素"


def test_interactive_preview_is_flashcard_only():
    from fastapi.testclient import TestClient
    from ai_engine.main import app

    node = {"title": "排序算法", "metadata": {"interactive_asset": {"items": [
        {"type": "quiz", "questions": [{
            "question": "测验题不应显示", "options": ["甲", "乙"], "answer": "A",
        }]},
        {"type": "flashcard", "front": "快速排序", "back": "分治思想"},
    ]}}}
    response = TestClient(app).post(
        "/api/v1/gen/interactive-preview", json={"course_name": "排序算法", "nodes": [node]},
    )
    assert response.status_code == 200
    assert response.json()["interactive_count"] == 1
    assert "快速排序" in response.json()["html"]
    assert "测验题不应显示" not in response.json()["html"]


def test_ppt_uses_flashcard_page_instead_of_stale_quiz():
    node = {
        "node_id": "card_1", "title": "排序算法", "type": "quiz",
        "interactive_plan": {"kind": "flashcard"},
        "metadata": {"interactive_asset": {"items": [
            {"type": "quiz", "questions": [{"question": "旧答题不应出现", "options": ["甲", "乙"], "answer": "A"}]},
            {"type": "flashcard", "front": "快速排序", "back": "分治思想"},
        ]}},
    }
    deck = Presentation(BytesIO(PPTExporter().export_to_bytes([node], "排序算法").getvalue()))
    texts = "\n".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert "快速排序" in texts and "分治思想" in texts
    assert "旧答题不应出现" not in texts
