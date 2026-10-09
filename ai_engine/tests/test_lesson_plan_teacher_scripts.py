"""教师讲稿须随 S3 版本节点进入标准 Word 教案。"""

from docx import Document
import pytest

from ai_engine.core.config import Settings
from ai_engine.gen import lesson_plan
from ai_engine.gen.export_doc import LessonPlanDOCExporter


NODES = [
    {
        "node_id": "n1",
        "title": "冒泡排序",
        "content": "幻灯片正文",
        "script": "先观察相邻元素。\n再演示交换过程。",
    },
    {
        "node_id": "n2",
        "title": "快速排序",
        "lecture_script": "通过基准元素讲解分区。",
    },
    {"node_id": "n3", "title": "练习", "content": "仅有课件正文"},
]


class StubLLM:
    def __init__(self, fail=False):
        self.fail = fail

    def chat_with_schema(self, *_args, **_kwargs):
        if self.fail:
            raise ConnectionError("offline test")
        return {
            "teaching_objectives": [{"dimension": "知识目标", "content": "掌握排序思想"}],
            "teaching_process": [{
                "phase": "新授", "duration_minutes": 35,
                "teacher_activity": "讲解算法", "student_activity": "观察演示",
                "design_intent": "理解算法",
            }],
        }


class ResponseLLM:
    def __init__(self, response):
        self.response = response
        self.calls = 0

    def chat_with_schema(self, *_args, **_kwargs):
        self.calls += 1
        if isinstance(self.response, list):
            return self.response[min(self.calls - 1, len(self.response) - 1)]
        return self.response


def test_provider_objective_map_and_nested_process_are_preserved(monkeypatch):
    monkeypatch.setattr(lesson_plan, "settings", Settings(GEN_MODE="llm"))
    llm = ResponseLLM({
        "teaching_theme": "排序算法",
        "teaching_objectives": {
            "knowledge": "理解比较排序",
            "ability": ["能手工追踪交换过程", "分析复杂度"],
            "literacy": "建立算法效率意识",
        },
        "teaching_process": [
            {"phase": "课堂导入", "duration_minutes": 5, "teacher_activity": "提出排序问题"},
            {"phase": "知识讲解", "duration_minutes": 30, "sub_phases": [
                {"teacher_activity": "演示冒泡排序", "student_activity": "追踪交换", "design_intent": "理解比较"},
                {"teacher_activity": "分析快速排序", "student_activity": "模拟分区"},
            ]},
            {"phase": "巩固练习", "duration_minutes": 5, "teacher_activity": "出示练习"},
            {"phase": "课堂小结", "duration_minutes": 5, "teacher_activity": "归纳复杂度"},
        ],
        "teaching_key_points": ["冒泡排序的交换"],
        "blackboard_design": "比较排序：冒泡与快速",
    })
    plan = lesson_plan.LessonPlanGenerator(llm_client=llm).generate(NODES, {"duration_minutes": 45})

    assert llm.calls == 1
    assert plan.data_source == "llm"
    assert plan.teaching_objectives[0] == {"dimension": "知识目标", "content": "理解比较排序"}
    assert "分析复杂度" in plan.teaching_objectives[1]["content"]
    assert "演示冒泡排序" in plan.teaching_process[1].teacher_activity
    assert "分析快速排序" in plan.teaching_process[1].teacher_activity
    assert plan.key_difficulties["key_points"] == ["冒泡排序的交换"]
    assert plan.board_design == "比较排序：冒泡与快速"


def test_empty_first_response_retries_small_core_schema(monkeypatch):
    monkeypatch.setattr(lesson_plan, "settings", Settings(GEN_MODE="llm"))
    llm = ResponseLLM([
        {"teaching_theme": "排序算法"},
        {"teaching_objectives": {"knowledge": "理解冒泡排序"}},
    ])
    plan = lesson_plan.LessonPlanGenerator(llm_client=llm).generate(NODES, {"topic": "排序算法"})

    assert llm.calls == 2
    assert plan.data_source == "llm"
    assert plan.teaching_objectives[0]["content"] == "理解冒泡排序"
    assert "教学过程" in plan.filled_sections


def test_missing_objectives_preserves_ai_process_and_reports_completion(monkeypatch):
    monkeypatch.setattr(lesson_plan, "settings", Settings(GEN_MODE="llm"))
    plan = lesson_plan.LessonPlanGenerator(llm_client=ResponseLLM({
        "teaching_process": [{
            "phase": "新授", "duration_minutes": 30,
            "teacher_activity": "AI 讲解冒泡排序的交换过程",
            "student_activity": "学生模拟相邻比较",
        }],
    })).generate(NODES, {"topic": "排序算法", "duration_minutes": 45})

    assert plan.data_source == "llm"
    assert "教学目标" in plan.filled_sections
    assert len(plan.teaching_objectives) == 3
    assert next(p for p in plan.teaching_process if p.phase == "新授").teacher_activity == "AI 讲解冒泡排序的交换过程"
    assert plan.to_dict()["filled_sections"] == plan.filled_sections


def test_missing_process_preserves_ai_objectives(monkeypatch):
    monkeypatch.setattr(lesson_plan, "settings", Settings(GEN_MODE="llm"))
    plan = lesson_plan.LessonPlanGenerator(llm_client=ResponseLLM({
        "teaching_objectives": [{"dimension": "知识目标", "content": "AI 制定的排序目标"}],
        "teaching_process": [{}],
    })).generate(NODES, {"topic": "排序算法"})

    assert plan.data_source == "llm"
    assert "教学过程" in plan.filled_sections
    assert plan.teaching_objectives[0]["content"] == "AI 制定的排序目标"
    assert len(plan.teaching_process) == 4


def test_no_usable_ai_sections_reports_safe_fallback_reason(monkeypatch):
    monkeypatch.setattr(lesson_plan, "settings", Settings(GEN_MODE="llm"))
    plan = lesson_plan.LessonPlanGenerator(llm_client=ResponseLLM({
        "teaching_process": [{"phase": "新授", "duration_minutes": 30}],
    })).generate(NODES, {"topic": "排序算法"})

    assert plan.data_source == "fallback"
    assert plan.fallback_reason == "AI 未返回可用的教学目标或教学过程"
    assert plan.to_dict()["fallback_reason"] == plan.fallback_reason


def test_provider_exception_is_not_exposed_as_fallback_reason(monkeypatch):
    monkeypatch.setattr(lesson_plan, "settings", Settings(GEN_MODE="llm"))

    class FailingLLM:
        def chat_with_schema(self, *_args, **_kwargs):
            raise RuntimeError("provider request failed with secret-token-placeholder")

    plan = lesson_plan.LessonPlanGenerator(llm_client=FailingLLM()).generate(NODES)

    assert plan.data_source == "fallback"
    assert plan.fallback_reason == "AI 教案生成异常，请查看 AI Engine 日志"
    assert "secret-token-placeholder" not in plan.to_dict()["fallback_reason"]


@pytest.mark.parametrize("fail", [False, True])
def test_teacher_scripts_are_copied_into_docx_for_llm_and_fallback(monkeypatch, fail):
    monkeypatch.setattr(lesson_plan, "settings", Settings(GEN_MODE="llm"))
    plan = lesson_plan.LessonPlanGenerator(llm_client=StubLLM(fail)).generate(
        NODES, {"topic": "排序算法", "duration_minutes": 45},
    )

    assert plan.data_source == ("fallback" if fail else "llm")
    assert [item["node_id"] for item in plan.teacher_scripts] == ["n1", "n2"]
    assert plan.teacher_scripts[0]["script"] == NODES[0]["script"]
    assert plan.to_dict()["teacher_scripts"] == plan.teacher_scripts

    doc = Document(LessonPlanDOCExporter().export_to_bytes(plan))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert text.index("四、教学过程") < text.index("教师讲稿（按课件节点）")
    assert text.index("教师讲稿（按课件节点）") < text.index("五、课堂活动设计")
    assert "先观察相邻元素。\n再演示交换过程。" in text
    assert "通过基准元素讲解分区。" in text
    assert "仅有课件正文" not in text


def test_lesson_plan_without_scripts_keeps_existing_sections(monkeypatch):
    monkeypatch.setattr(lesson_plan, "settings", Settings(GEN_MODE="llm"))
    plan = lesson_plan.LessonPlanGenerator(llm_client=StubLLM()).generate(
        [{"node_id": "n1", "title": "排序算法", "content": "课件正文"}],
        {"topic": "排序算法"},
    )
    doc = Document(LessonPlanDOCExporter().export_to_bytes(plan))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    assert "教师讲稿（按课件节点）" not in text
    assert "四、教学过程" in text
    assert "五、课堂活动设计" in text
