"""Structural and conservative layout checks for exported PPTX artifacts."""

from __future__ import annotations

import io
import math
import zipfile
from typing import Any

from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE


def _finding(code: str, severity: str, message: str, slide_number: int | None = None) -> dict[str, Any]:
    item: dict[str, Any] = {"code": code, "severity": severity, "message": message}
    if slide_number is not None:
        item["slide_number"] = slide_number
    return item


def validate_pptx(data: bytes) -> dict[str, Any]:
    """Inspect a completed PPTX without modifying its bytes.

    Text-fit detection is a conservative geometry heuristic, not a rendering
    guarantee. Geometry warnings on non-text shapes allow intentional bleeds.
    """
    findings: list[dict[str, Any]] = []
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            corrupt_member = archive.testzip()
            if corrupt_member:
                findings.append(_finding(
                    "invalid_package", "error", "PPTX package contains a damaged component."
                ))
                return _report(findings, 0)
        presentation = Presentation(io.BytesIO(data))
    except Exception:
        findings.append(_finding(
            "invalid_presentation", "error", "PPTX package could not be opened as a presentation."
        ))
        return _report(findings, 0)

    slide_count = len(presentation.slides)
    if not slide_count:
        findings.append(_finding(
            "empty_presentation", "error", "PPTX contains no slides."
        ))
        return _report(findings, slide_count)

    slide_width = presentation.slide_width
    slide_height = presentation.slide_height
    edge_tolerance = 914400 * 0.02
    for slide_number, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            left = shape.left
            top = shape.top
            width = shape.width
            height = shape.height
            if None in (left, top, width, height):
                continue
            outside = (
                left < -edge_tolerance or top < -edge_tolerance
                or left + width > slide_width + edge_tolerance
                or top + height > slide_height + edge_tolerance
            )
            has_text = bool(getattr(shape, "has_text_frame", False) and shape.text.strip())
            if outside:
                findings.append(_finding(
                    "text_outside_slide" if has_text else "shape_outside_slide",
                    "error" if has_text else "warning",
                    "Text content extends beyond the slide canvas."
                    if has_text else "A shape extends beyond the slide canvas.",
                    slide_number,
                ))
                continue
            if has_text and _may_overflow_text(shape):
                findings.append(_finding(
                    "text_may_overflow", "warning",
                    "Text may exceed its fixed text-box capacity; visual rendering was not checked.",
                    slide_number,
                ))

    return _report(findings, slide_count)


def _may_overflow_text(shape: Any) -> bool:
    frame = shape.text_frame
    if frame.auto_size in {
        MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT,
        MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE,
    }:
        return False
    box_width = max(shape.width, 1) / 914400
    box_height = max(shape.height, 1) / 914400
    text = frame.text.strip()
    if not text or box_width <= 0 or box_height <= 0:
        return False

    font_sizes = [
        run.font.size.pt
        for paragraph in frame.paragraphs
        for run in paragraph.runs
        if run.font.size is not None
    ]
    font_size = max(font_sizes, default=18)
    char_width = font_size / 72 * 0.95
    line_height = font_size / 72 * 1.25
    usable_width = max(box_width - 0.12, 0.1)
    usable_height = max(box_height - 0.08, 0.05)
    chars_per_line = max(int(usable_width / char_width), 1)
    lines_available = max(int(usable_height / line_height), 1)
    estimated_lines = sum(
        max(math.ceil(len(line) / chars_per_line), 1)
        for paragraph in frame.paragraphs
        for line in (paragraph.text.splitlines() or [""])
    )
    return estimated_lines > lines_available


def _report(findings: list[dict[str, Any]], slide_count: int) -> dict[str, Any]:
    if any(item["severity"] == "error" for item in findings):
        status = "failed"
    elif findings:
        status = "passed_with_warnings"
    else:
        status = "passed"
    return {"status": status, "slide_count": slide_count, "findings": findings}
