"""S3 density must affect slide content in both preview and editable PPTX."""

from pptx import Presentation
from pptx.util import Inches

from ai_engine.gen.export_ppt import PPTExporter
from ai_engine.gen.presentation import content_geometry, density_profile
from ai_engine.gen.preview import PreviewRenderer


CONTENT = "快速排序采用分治思想"
POINT = "选择基准并划分数组"


def _node(density):
    return {
        "node_id": "density-content",
        "type": "content",
        "title": "排序算法",
        "content": CONTENT,
        "key_points": [POINT],
        "metadata": {"presentation": {"density": density}},
    }


def _text_size(prs, text):
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text == text:
                return shape.text_frame.paragraphs[0].runs[0].font.size.pt
    raise AssertionError(f"Missing slide text: {text}")


def test_density_changes_visible_preview_typography():
    for name in ("compact", "comfortable", "minimal"):
        style = density_profile(name)
        html = PreviewRenderer().render([_node(name)], "排序算法")
        assert f'data-density="{name}"' in html
        assert f".presentation .svg-block .body {{ font-size: {style.preview_body_px}px; }}" in html
        assert f".presentation .svg-block .key-points {{ gap: {style.card_gap_px}px; }}" in html
        assert CONTENT in html and POINT in html


def test_density_changes_exported_editable_text_without_dropping_content():
    body_sizes = []
    point_sizes = []
    for name in ("compact", "comfortable", "minimal"):
        prs = Presentation(PPTExporter().export_to_bytes([_node(name)], "排序算法"))
        body_sizes.append(_text_size(prs, CONTENT))
        point_sizes.append(_text_size(prs, f"· {POINT}"))
        assert all(shape.shape_type != 13 for slide in prs.slides for shape in slide.shapes)
    assert body_sizes == [13, 15, 17]
    assert point_sizes[0] < point_sizes[1] < point_sizes[2]


def test_density_moves_content_panels_in_preview_and_pptx():
    panel_positions = []
    for name in ("compact", "comfortable", "minimal"):
        geometry = content_geometry(name)
        html = PreviewRenderer().render([_node(name)], "排序算法")
        expected_svg_x = geometry.left[0] * 96
        assert f'x="{expected_svg_x:.1f}"' in html

        prs = Presentation(PPTExporter().export_to_bytes([_node(name)], "排序算法"))
        content_slide = next(slide for slide in prs.slides
                             if any(getattr(shape, "text", "") == CONTENT for shape in slide.shapes))
        assert any(abs(shape.left - Inches(geometry.left[0])) < Inches(.01)
                   and abs(shape.top - Inches(geometry.top)) < Inches(.01)
                   and abs(shape.width - Inches(geometry.left[1])) < Inches(.01)
                   for shape in content_slide.shapes)
        panel_positions.append(geometry.left[0])
    assert len(set(panel_positions)) == 3


def test_unknown_density_keeps_comfortable_defaults():
    assert density_profile("unknown") == density_profile("comfortable")
