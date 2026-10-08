"""H8 互动 HTML5 集成进 PPT —— 互动页占位形状 + 超链接指向同目录 html5 文件。

验收点：
  1. 提供 html5 文件名时，quiz/interactive 页导出的 PPTX 带一个可点击形状，
     超链接地址 = 同目录相对文件名；
  2. 不提供 html5 文件名时（bytes 直出等旧路径）不产生死链接；
  3. 内容页 / 封面等非互动页不挂超链接。
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

_ROOT = Path(__file__).resolve().parent.parent.parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from ai_engine.gen.export_ppt import PPTExporter  # noqa: E402


def _hyperlink_addresses(prs) -> set[str]:
    addresses = set()
    for slide in prs.slides:
        for shape in slide.shapes:
            try:
                address = shape.click_action.hyperlink.address
            except (AttributeError, Exception):
                continue
            if address:
                addresses.add(str(address))
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run_address = getattr(run.hyperlink, "address", None)
                        if run_address:
                            addresses.add(str(run_address))
    return addresses


def _quiz_node() -> dict:
    return {
        "node_id": "n_quiz",
        "type": "quiz",
        "title": "课堂小测",
        "questions": [{"question": "冒泡排序的时间复杂度是？", "options": ["O(n)", "O(n²)"],
                       "answer": "B", "explanation": "双重循环。"}],
    }


def test_interactive_slide_links_to_sibling_html5():
    buf = PPTExporter().export_to_bytes([_quiz_node()], "互动课件", html5_filename="exp_demo.html")

    from pptx import Presentation

    prs = Presentation(io.BytesIO(buf.getvalue()))
    addresses = _hyperlink_addresses(prs)
    assert "exp_demo.html" in addresses, "互动页应带指向同目录 html5 的超链接"


def test_interactive_slide_without_html5_has_no_dead_link():
    buf = PPTExporter().export_to_bytes([_quiz_node()], "互动课件")

    from pptx import Presentation

    prs = Presentation(io.BytesIO(buf.getvalue()))
    assert _hyperlink_addresses(prs) == set(), "未提供 html5 文件时不应产生死链接"


def test_content_slide_never_links_to_html5():
    node = {
        "node_id": "n1", "type": "content", "title": "普通内容页",
        "content": "普通内容，不应出现互动超链接。",
        "metadata": {},
    }
    buf = PPTExporter().export_to_bytes([node], "普通课件", html5_filename="exp_demo.html")

    from pptx import Presentation

    prs = Presentation(io.BytesIO(buf.getvalue()))
    assert _hyperlink_addresses(prs) == set(), "非互动页不应挂 html5 超链接"
