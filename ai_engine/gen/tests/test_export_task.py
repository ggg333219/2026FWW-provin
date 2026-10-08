"""导出模块单元测试 — test_export_task.py

测试目标：
  PPTExporter / DOCExporter 导出输出正确格式的文件。
  各种节点类型覆盖。
  异常情况（空节点、非串行化内容）。
"""

import os
import json
import zipfile
import unittest
from pathlib import Path
from io import BytesIO

from pptx import Presentation

from ai_engine.gen.export_ppt import PPTExporter, ExportPPTConfig, resolve_node_type
from ai_engine.gen.export_doc import DOCExporter, ExportDocConfig

# =============================================================================
# 公共测试夹具
# =============================================================================

def _content_nodes() -> list[dict]:
    """生成一组典型的教学节点"""
    return [
        {"node_id": "n1", "title": "课程封面", "type": "title",
         "content": "人工智能导论", "summary": "AI入门课程"},
        {"node_id": "n2", "title": "什么是AI", "type": "content",
         "content": "人工智能（AI）是模拟人类智能的技术。\n包括机器学习、深度学习等。",
         "core_points": ["定义", "范畴"], "summary": "AI基本概念"},
        {"node_id": "n3", "title": "课堂测验", "type": "quiz",
         "questions": ["AI的中文全称是？", "以下哪个不是AI技术？"],
         "summary": "检验理解"},
        {"node_id": "n4", "title": "本章小结", "type": "summary",
         "key_points": ["AI=模拟人类智能", "三大技术分支"],
         "summary": "核心回顾"},
    ]


def _empty_content_nodes() -> list[dict]:
    """带空字段的节点"""
    return [
        {"node_id": "n1", "title": "", "type": "content", "content": ""},
    ]


def _minimal_nodes() -> list[dict]:
    """最小节点"""
    return [{"node_id": "n1", "title": "测试"}]


def _node_with_segments() -> list[dict]:
    """带 segments 的节点（增强模式输出）"""
    return [
        {"node_id": "n1", "title": "开场", "type": "content", "content": "欢迎",
         "segments": [
             {"segment_id": "seg_1", "text": "欢迎来到AI课堂", "segment_type": "opening"},
             {"segment_id": "seg_2", "text": "这是主要内容", "segment_type": "explanation"},
         ]},
    ]


# =============================================================================
# PPTExporter 测试
# =============================================================================

class TestPPTExporter(unittest.TestCase):
    """PPTX 导出功能测试"""

    def setUp(self):
        self.ppt = PPTExporter(ExportPPTConfig(output_dir="/tmp"))

    # ── 基础功能 ──

    def test_export_minimal_nodes(self):
        """最少节点也能导出有效 PPTX"""
        result = self.ppt.export(_minimal_nodes(), "最小测试")
        self.assertGreater(result["pages"], 0)
        self.assertGreater(result["file_size"], 0)
        self.assertTrue(os.path.exists(result["path"]))

    def test_export_content_nodes(self):
        """典型教学节点导出为多页 PPTX"""
        result = self.ppt.export(_content_nodes(), "内容测试")
        self.assertGreaterEqual(result["pages"], 3)
        with zipfile.ZipFile(result["path"]) as z:
            self.assertIn("[Content_Types].xml", z.namelist())
            slide_files = [n for n in z.namelist() if n.startswith("ppt/slides/slide")]
            self.assertGreaterEqual(len(slide_files), 2)

    def test_export_empty_content_still_creates_slides(self):
        """空内容节点仍生成有效幻灯片"""
        result = self.ppt.export(_empty_content_nodes(), "空内容")
        self.assertGreater(result["pages"], 0)
        self.assertGreater(result["file_size"], 100)

    # ── export_to_bytes ──

    def test_export_to_bytes_returns_valid_zip(self):
        """字节流输出是有效 ZIP"""
        buf = self.ppt.export_to_bytes(_content_nodes(), "字节测试")
        self.assertGreater(buf.getbuffer().nbytes, 0)
        with zipfile.ZipFile(buf) as z:
            self.assertIn("[Content_Types].xml", z.namelist())

    def test_export_to_bytes_content_types(self):
        """字节流包含幻灯片内容"""
        buf = self.ppt.export_to_bytes(_content_nodes(), "内容检查")
        with zipfile.ZipFile(buf) as z:
            slides = [n for n in z.namelist() if n.startswith("ppt/slides/slide")]
            self.assertGreaterEqual(len(slides), 2)

    def test_export_to_bytes_minimal(self):
        """最少节点字节流"""
        buf = self.ppt.export_to_bytes(_minimal_nodes(), "最小字节")
        self.assertGreater(buf.getbuffer().nbytes, 100)

    def test_export_to_bytes_segments(self):
        """segments 的类型、文本与时长均进入可编辑文本框"""
        nodes = _node_with_segments()
        nodes[0]["segments"][0]["estimated_seconds"] = 45
        buf = self.ppt.export_to_bytes(nodes, "分段节点")
        text = "\n".join(
            shape.text
            for slide in Presentation(buf).slides
            for shape in slide.shapes
            if hasattr(shape, "text")
        )
        self.assertIn("导入 · 45秒", text)
        self.assertIn("欢迎来到AI课堂", text)
        self.assertIn("讲解", text)
        self.assertIn("这是主要内容", text)

    def test_type_precedes_node_type_like_preview(self):
        """与预览侧保持一致：冲突时 type 优先于 node_type。"""
        nodes = [{
            "node_id": "n1", "title": "课堂互动", "node_type": "cover",
            "type": "quiz", "questions": ["哪个字段优先？"],
        }]
        buf = self.ppt.export_to_bytes(nodes, "双字段")
        prs = Presentation(buf)
        self.assertEqual(len(prs.slides), 1)
        slide_text = "\n".join(shape.text for shape in prs.slides[0].shapes if hasattr(shape, "text"))
        # 标题条只剩标题与页码：kicker（"INTERACTION"）与标题下的强调色带是导出侧
        # 独有、预览任何版式里都没有的装饰，已去掉。这里改用页面标题断言页型判定。
        self.assertIn("课堂互动", slide_text)
        self.assertIn("哪个字段优先？", slide_text)

    def test_interactive_node_exports_explicit_front_and_back(self):
        nodes = [{
            "node_id": "flip_1",
            "type": "interactive",
            "node_type": "interactive",
            "title": "定积分知识翻卡",
            "content": "正面：定积分的几何意义是什么？\n背面：曲边梯形面积的代数和。",
        }]

        prs = Presentation(self.ppt.export_to_bytes(nodes, "定积分知识翻卡"))
        text = "\n".join(
            shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text")
        )

        self.assertIn("定积分的几何意义是什么？", text)
        self.assertIn("曲边梯形面积的代数和。", text)
        self.assertNotIn("请结合本节内容提出一个关键问题。", text)

    def test_key_points_and_core_points_are_both_exported(self):
        """新旧要点字段同时存在时合并导出，不再产生空白要点页。"""
        nodes = [{
            "node_id": "n1", "title": "总结", "node_type": "summary",
            "key_points": ["关键点A"], "core_points": ["核心点B", "关键点A"],
        }]
        buf = self.ppt.export_to_bytes(nodes, "要点兼容")
        text = "\n".join(
            shape.text
            for slide in Presentation(buf).slides
            for shape in slide.shapes
            if hasattr(shape, "text")
        )
        self.assertEqual(text.count("关键点A"), 1)
        self.assertIn("核心点B", text)

    def test_three_teaching_templates_keep_native_editable_text(self):
        """三套模板均生成 16:9、无整页图片伪导出的原生文本。"""
        for template_id in ("default", "academic", "minimal"):
            with self.subTest(template_id=template_id):
                buf = self.ppt.export_to_bytes(_content_nodes(), "模板测试", template_id)
                prs = Presentation(buf)
                self.assertAlmostEqual(prs.slide_width / prs.slide_height, 16 / 9, places=2)
                texts = [
                    shape.text for slide in prs.slides for shape in slide.shapes
                    if hasattr(shape, "text") and shape.text.strip()
                ]
                self.assertTrue(any("什么是AI" in text for text in texts))
                self.assertFalse(any(shape.shape_type == 13 for slide in prs.slides for shape in slide.shapes))

    def test_template_file_samples_are_not_leaked_into_export(self):
        """从模板目录加载时会清除样例页，不污染节点同源页数。"""
        template_dir = Path(__file__).resolve().parents[1] / "templates"
        exporter = PPTExporter(ExportPPTConfig(template_dir=str(template_dir), output_dir="/tmp"))
        prs = Presentation(exporter.export_to_bytes(
            [{"node_id": "n1", "node_type": "slide", "title": "唯一节点", "content": "唯一正文"}],
            "母版加载", "academic",
        ))
        self.assertEqual(len(prs.slides), 1)
        text = "\n".join(shape.text for shape in prs.slides[0].shapes if hasattr(shape, "text"))
        self.assertIn("唯一节点", text)
        self.assertNotIn("课程目录", text)

    def test_long_content_uses_continuation_slides_without_truncation(self):
        """超长正文自动续页，首尾内容均保留。"""
        long_text = "起始标记。" + "这是用于验证自适应分页的完整讲授正文。" * 80 + "结束标记。"
        nodes = [{"node_id": "n1", "node_type": "slide", "title": "长内容", "content": long_text}]
        prs = Presentation(self.ppt.export_to_bytes(nodes, "长内容测试"))
        self.assertGreater(len(prs.slides), 1)
        text = "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text"))
        self.assertIn("起始标记", text)
        self.assertIn("结束标记", text)
        # 续页标题沿用本页标题 +「（续 N）」（export_ppt._add_content_slide）。改动前
        # 分页走的是通用兜底文案「正文（续）」，与页面标题脱节；现在内容页自己分页，
        # 续页仍属于同一页内容，标题就该还是原标题。
        self.assertIn("长内容（续 2）", text)

    def test_body_panel_pagination_matches_what_it_renders(self):
        """右下正文框：分页算的高度必须就是渲染要画的高度。

        改动前这里是两套算法 —— 分页按「字符数 ÷ 每行字数 + 1」估权重（给每个条目
        多算整整一行当段距，实际只有 0.1in），渲染按「行数 × 行高 + 段距」排，于是
        这块 2.76in 的面板只认 3 条要点，第 4 条被推到下一页，而面板里还剩一半高度。
        这正是「右下角的框总有一行到下一页去」的成因。
        """
        node = {
            "node_id": "n-fit", "node_type": "content", "title": "分页容量",
            "point_details": [
                {"point": f"要点{i}", "detail": f"展开{i}"} for i in range(1, 6)
            ],
            "figure": {"kind": "flow", "items": ["甲", "乙", "丙"]},
        }
        prs = Presentation(self.ppt.export_to_bytes([node], "分页容量"))
        text = "\n".join(
            shape.text for slide in prs.slides for shape in slide.shapes
            if hasattr(shape, "text")
        )

        # 五条一行要点（每条约 0.36in）放得进这块面板，就该在同一页 ——
        # 更不能出现"分页塞进来了、渲染的 cursor >= limit 又把它丢掉"。
        self.assertEqual(len(prs.slides), 1)
        for i in range(1, 6):
            self.assertIn(f"要点{i}：展开{i}", text)

    def test_two_column_uses_independent_left_and_right_panels(self):
        """two_column 不再收敛为普通内容页，左右栏保持固定位置。"""
        node = {
            "node_id": "n-two", "node_type": "two_column", "title": "双栏教学",
            "key_points": ["左栏要点A", "左栏要点B"],
            "content": "右栏正文回退",
            "segments": [{
                "segment_type": "explanation", "text": "右栏讲授内容", "estimated_seconds": 60,
            }],
        }
        self.assertEqual(resolve_node_type(node), "two_column")
        prs = Presentation(self.ppt.export_to_bytes([node], "双栏验证"))
        self.assertEqual(len(prs.slides), 1)

        left_shapes = [
            shape for shape in prs.slides[0].shapes
            if hasattr(shape, "text") and "左栏要点A" in shape.text
        ]
        right_shapes = [
            shape for shape in prs.slides[0].shapes
            if hasattr(shape, "text") and "右栏讲授内容" in shape.text
        ]
        self.assertEqual(len(left_shapes), 1)
        self.assertEqual(len(right_shapes), 1)
        self.assertLess(left_shapes[0].left.inches, 6.0)
        self.assertGreater(right_shapes[0].left.inches, 6.0)
        all_text = "\n".join(shape.text for shape in prs.slides[0].shapes if hasattr(shape, "text"))
        self.assertIn("双栏教学", all_text)
        self.assertIn("讲解 · 1分钟", all_text)

    # ── 模板 ──

    def test_export_with_default_template(self):
        """默认模板 id = 'default'"""
        result = self.ppt.export(_content_nodes(), "默认模板", template_id="default")
        self.assertGreater(result["pages"], 0)

    # ── 文件名 ──

    def test_export_filename_contains_course_name(self):
        """输出文件名包含课程名"""
        result = self.ppt.export(_content_nodes(), "文件名测试")
        self.assertIn("文件名测试", result["path"])

    # ── 批量 / 多种节点混合 ──

    def test_export_mixed_types(self):
        """title + content + quiz + summary 混合"""
        nodes = [
            {"node_id": "n1", "title": "封面", "type": "title", "content": "标题"},
            {"node_id": "n2", "title": "正文", "type": "content", "content": "内容", "core_points": ["A"]},
            {"node_id": "n3", "title": "问答", "type": "quiz", "questions": ["Q1"]},
            {"node_id": "n4", "title": "总结", "type": "summary", "key_points": ["K1"]},
        ]
        result = self.ppt.export(nodes, "混合")
        self.assertGreaterEqual(result["pages"], 3)

    def test_export_large_number_of_nodes(self):
        """20个节点快速导出不报错"""
        nodes = [{"node_id": f"n{i}", "title": f"第{i}节", "type": "content", "content": f"这是第{i}节内容"}
                 for i in range(20)]
        result = self.ppt.export(nodes, "大批量")
        self.assertGreaterEqual(result["pages"], 19)

    def test_export_no_type_field(self):
        """节点没有 type 字段也正常处理"""
        nodes = [{"node_id": "n1", "title": "第一节", "content": "内容"}]
        buf = self.ppt.export_to_bytes(nodes, "无类型")
        self.assertGreater(buf.getbuffer().nbytes, 0)

    # ── 一致性 ──

    def test_export_idempotent(self):
        """相同输入两次导出，都生成有效文件"""
        nodes = _content_nodes()
        r1 = self.ppt.export(nodes, "幂等")
        r2 = self.ppt.export(nodes, "幂等")
        self.assertGreater(r1["pages"], 0)
        self.assertGreater(r2["pages"], 0)

    def test_export_no_course_name(self):
        """course_name 为空时不崩溃"""
        nodes = _minimal_nodes()
        result = self.ppt.export(nodes, "")
        self.assertGreater(result["pages"], 0)


# =============================================================================
# DOCExporter 测试
# =============================================================================

class TestDOCExporter(unittest.TestCase):
    """DOCX 导出功能测试"""

    def setUp(self):
        self.doc = DOCExporter(ExportDocConfig(output_dir="/tmp"))

    # ── 基础功能 ──

    def test_export_minimal_nodes(self):
        """最少节点也能导出有效 DOCX"""
        buf = self.doc.export_to_bytes(_minimal_nodes(), "最小测试")
        self.assertGreater(buf.getbuffer().nbytes, 0)
        with zipfile.ZipFile(buf) as z:
            self.assertIn("word/document.xml", z.namelist())

    def test_export_content_nodes(self):
        """典型教学节点导出 DOCX"""
        result = self.doc.export(_content_nodes(), "内容测试")
        self.assertGreater(result["sections"], 0)
        self.assertGreater(result["file_size"], 0)

    def test_export_empty_content_still_creates_doc(self):
        """空内容节点仍生成有效文档"""
        result = self.doc.export(_empty_content_nodes(), "空内容")
        self.assertGreater(result["sections"], 0)

    # ── export_to_bytes ──

    def test_export_to_bytes_returns_valid_docx_zip(self):
        """字节流输出是有效 DOCX ZIP"""
        buf = self.doc.export_to_bytes(_content_nodes(), "字节测试")
        self.assertGreater(buf.getbuffer().nbytes, 0)
        with zipfile.ZipFile(buf) as z:
            self.assertIn("word/document.xml", z.namelist())

    def test_export_to_bytes_minimal(self):
        """最少节点字节流"""
        buf = self.doc.export_to_bytes(_minimal_nodes(), "最小字节")
        self.assertGreater(buf.getbuffer().nbytes, 100)

    def test_export_to_bytes_segments(self):
        """带 segments 的节点正常输出"""
        buf = self.doc.export_to_bytes(_node_with_segments(), "分段节点")
        self.assertGreater(buf.getbuffer().nbytes, 0)
        with zipfile.ZipFile(buf) as z:
            self.assertIn("word/document.xml", z.namelist())

    # ── 混合类型 ──

    def test_export_mixed_types(self):
        """title + content + quiz + summary 混合"""
        nodes = [
            {"node_id": "n1", "title": "封面", "type": "title", "content": "标题"},
            {"node_id": "n2", "title": "正文", "type": "content", "content": "内容"},
            {"node_id": "n3", "title": "问答", "type": "quiz", "questions": ["Q1"]},
            {"node_id": "n4", "title": "总结", "type": "summary", "key_points": ["K1"]},
        ]
        result = self.doc.export(nodes, "混合")
        self.assertGreater(result["sections"], 0)

    def test_export_large_number_of_nodes(self):
        """20个节点快速导出不报错"""
        nodes = [{"node_id": f"n{i}", "title": f"第{i}节", "type": "content", "content": f"这是第{i}节内容"}
                 for i in range(20)]
        result = self.doc.export(nodes, "大批量")
        self.assertGreater(result["sections"], 0)

    def test_export_no_type_field(self):
        """节点没有 type 字段也正常处理"""
        nodes = [{"node_id": "n1", "title": "第一节", "content": "内容"}]
        buf = self.doc.export_to_bytes(nodes, "无类型")
        self.assertGreater(buf.getbuffer().nbytes, 0)

    def test_export_idempotent(self):
        """相同输入两次导出，都生成有效文件"""
        nodes = _content_nodes()
        r1 = self.doc.export(nodes, "幂等")
        r2 = self.doc.export(nodes, "幂等")
        self.assertGreater(r1["sections"], 0)
        self.assertGreater(r2["sections"], 0)


# =============================================================================
# 交叉格式一致性测试
# =============================================================================

class TestExportConsistency(unittest.TestCase):
    """同一组节点同时导出 PPTX 和 DOCX，均有效"""

    def test_pptx_and_docx_both_valid(self):
        """两种格式同时导出，内容一致"""
        nodes = _content_nodes()
        ppt = PPTExporter(ExportPPTConfig(output_dir="/tmp"))
        doc = DOCExporter(ExportDocConfig(output_dir="/tmp"))

        ppt_result = ppt.export(nodes, "交叉验证")
        doc_result = doc.export(nodes, "交叉验证")
        doc_buf = doc.export_to_bytes(nodes, "交叉验证")

        self.assertGreater(ppt_result["pages"], 0)
        self.assertGreater(doc_result["sections"], 0)
        self.assertGreater(doc_buf.getbuffer().nbytes, ppt_result["file_size"] // 2)

    def test_pptx_export_to_bytes_is_deterministic_format(self):
        """export_to_bytes 输出始终是 application/vnd.openxmlformats PPTX"""
        nodes = _minimal_nodes()
        ppt = PPTExporter(ExportPPTConfig(output_dir="/tmp"))
        buf = ppt.export_to_bytes(nodes, "格式检查")
        # 第一个字节是 PK (ZIP 签名)
        self.assertEqual(buf.getvalue()[:2], b"PK")


# =============================================================================
# 导出任务生命周期测试（不依赖 main.py 端点，直接使用 exporter 对象）
# =============================================================================

class TestExportLifecycle(unittest.TestCase):
    """模拟导出任务的生命周期"""

    def test_export_then_read_back_file(self):
        """导出后文件可读"""
        ppt = PPTExporter(ExportPPTConfig(output_dir="/tmp"))
        result = ppt.export(_content_nodes(), "生命周期测试")
        self.assertTrue(os.path.exists(result["path"]))
        self.assertGreater(os.path.getsize(result["path"]), 0)

    def test_export_to_bytes_then_save(self):
        """字节流保存为文件后有效"""
        ppt = PPTExporter(ExportPPTConfig(output_dir="/tmp"))
        buf = ppt.export_to_bytes(_content_nodes(), "字节保存")
        tmp = Path("/tmp") / f"_test_export_bytes_{hash(str(_content_nodes()))}.pptx"
        tmp.write_bytes(buf.getvalue())
        self.assertTrue(tmp.exists())
        self.assertGreater(tmp.stat().st_size, 0)
        # 验证是有效 ZIP
        with zipfile.ZipFile(str(tmp)) as z:
            self.assertIn("[Content_Types].xml", z.namelist())
        tmp.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main(verbosity=2)
