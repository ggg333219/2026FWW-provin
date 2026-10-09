"""The preview SVG and editable PPTX should share the same slide canvas and page styles."""

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ai_engine.gen.export_ppt import ExportPPTConfig, PPTExporter
from ai_engine.gen.page_spec import build_page_spec
from ai_engine.gen.presentation import theme_from_palette
from ai_engine.gen.preview import PreviewRenderer
from ai_engine.gen.svg_renderer import render_page_spec_to_svg


def _ppt(node):
    config = ExportPPTConfig(
        include_auto_cover=False, include_auto_toc=False,
        include_auto_summary=False, include_closing_slide=False,
    )
    return Presentation(PPTExporter(config).export_to_bytes([node]))


def test_preview_canvas_is_full_bleed_and_has_no_extra_slide_badge():
    html = PreviewRenderer().render_node({"type": "cover", "title": "排序算法"})
    assert 'viewBox="0 0 1280 720"' in html
    page = PreviewRenderer().render([{"type": "cover", "title": "排序算法"}])
    assert "padding: 0; border: 0; border-radius: 0; aspect-ratio: 16 / 9" in page
    assert ".slide-num { display: none; }" in page


def test_light_palette_uses_the_same_dark_inverse_text_in_preview_and_ppt():
    node = {"type": "section", "title": "算法基础",
            "metadata": {"presentation": {"palette": "lemon-yellow"}}}
    html = PreviewRenderer().render([node])
    theme = theme_from_palette("lemon-yellow")
    assert f"--preview-inverse: #{theme.inverse}" in html
    prs = _ppt(node)
    title = next(shape for shape in prs.slides[0].shapes
                 if shape.has_text_frame and shape.text == "算法基础")
    assert str(title.text_frame.paragraphs[0].font.color.rgb) == theme.inverse


def test_agenda_uses_its_own_layout_and_exports_same_items():
    node = {"node_type": "agenda", "title": "目录", "key_points": ["冒泡排序", "快速排序"]}
    html = PreviewRenderer().render_node(node)
    assert 'class="slide agenda-slide"' in html
    assert 'class="section-list"' in html
    assert "冒泡排序" in html and "快速排序" in html
    prs = _ppt(node)
    text = "\n".join(shape.text for shape in prs.slides[0].shapes if shape.has_text_frame)
    assert "冒泡排序" in text and "快速排序" in text


def test_section_background_and_points_match_preview_theme():
    node = {"type": "section", "title": "算法基础", "key_points": ["比较排序", "分治思想"]}
    theme = theme_from_palette("classic")
    svg = render_page_spec_to_svg(build_page_spec(node, 1), theme)
    assert 'fill="#' + theme.secondary + '"' in svg
    assert 'fill="#' + theme.inverse + '"' in svg
    prs = _ppt(node)
    slide = prs.slides[0]
    assert str(slide.background.fill.fore_color.rgb) == theme.primary
    text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
    assert "比较排序" in text and "分治思想" in text


def test_cover_has_no_export_only_circles():
    prs = _ppt({"type": "cover", "title": "排序算法"})
    assert not any(shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and
                   shape.auto_shape_type == 9  # OVAL
                   for shape in prs.slides[0].shapes)


def test_quiz_and_summary_include_their_preview_panels():
    quiz = _ppt({"type": "quiz", "title": "课堂检测", "questions": ["什么是分治？"]})
    assert any(abs(shape.left / 914400 - .72) < .01 and
               abs(shape.top / 914400 - 1.42) < .01
               for shape in quiz.slides[0].shapes)
    summary = _ppt({"type": "summary", "title": "本节小结", "key_points": ["分治思想"]})
    assert "本节内容结束" in "\n".join(
        shape.text for shape in summary.slides[0].shapes if shape.has_text_frame
    )
