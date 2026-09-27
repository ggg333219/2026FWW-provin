"""P3：导出侧消费 page_spec 的图示与结论，且与预览同源。

这一组补的是改动前**完全没有覆盖**的部分：六个图示渲染器（flow 在 P3 之前就有，
另外五个是新的）此前没有任何测试碰过 —— 导出时抛异常、或者悄悄画出一张图片，
都不会有测试报警。

三件事必须同时锁住：

  1. 六种 kind 都能画出来，且**不产生 PICTURE**（契约 #1：.pptx 必须可编辑）。
     改成贴图最省事，但教师就再也改不动自己课件里的一张流程图了。
  2. 图示标签与大字结论**两端都在** —— 这正是本轮的立项理由：P1 的 figure /
     key_conclusion 只进了预览，学生拿到的 .pptx 里没有。
  3. 正文与要点并存时要点不能丢。page_spec 会把 bullets 降为 ROLE_SECONDARY，
     导出侧若只认 PRIMARY，要点会整块消失而 error 为空。
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from pptx import Presentation

_ROOT = Path(__file__).resolve().parent.parent.parent.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from ai_engine.gen.export_ppt import PPTExporter  # noqa: E402
from ai_engine.gen.preview import PreviewRenderer  # noqa: E402

FIGURE_KINDS = ["flow", "compare", "cycle", "hierarchy", "chart", "matrix"]

#: 每种 kind 都给一份**模型真会给的形状**（不是最干净的理想输入）。
_FIGURES: dict[str, dict] = {
    "flow": {"kind": "flow", "items": ["读取输入", "计算规模", "输出结果"]},
    "compare": {
        "kind": "compare",
        "left": {"title": "冒泡排序", "items": ["稳定", "平均 O(n²)"]},
        "right": {"title": "快速排序", "items": ["不稳定", "平均 O(n log n)"]},
    },
    "cycle": {"kind": "cycle", "items": ["采集", "处理", "反馈", "调整"]},
    "hierarchy": {"kind": "hierarchy", "items": ["排序算法", "冒泡排序", "快速排序"]},
    "chart": {
        "kind": "chart",
        "series": [{"label": "冒泡排序", "value": 45}, {"label": "快速排序", "value": 12}],
    },
    "matrix": {
        "kind": "matrix",
        "columns": ["算法", "稳定性", "平均复杂度"],
        "rows": [["冒泡排序", "稳定", "O(n²)"], ["快速排序", "不稳定", "O(n log n)"]],
    },
}

CONCLUSION = "复杂度描述的是增长趋势"


def _figure_node(kind: str) -> dict:
    return {
        "node_id": f"fig_{kind}", "type": "content", "node_type": "content",
        "title": f"{kind} 图示页", "content": "图示与结论必须两端同源。",
        "key_conclusion": CONCLUSION,
        "figure": _FIGURES[kind],
    }


def _ppt_text(nodes: list[dict]) -> str:
    prs = Presentation(PPTExporter().export_to_bytes(nodes, "图示验收"))
    return "\n".join(
        shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text")
    )


def _figure_tokens(figure: dict) -> list[str]:
    """图示里所有**由作者/模型写下**的字符串 —— 这些两端必须逐字一致。

    刻意不含 chart 的数值：预览侧用 diagrams._format_number 格式化，导出侧用
    f"{v:g}"，常见取值下一致，但没必要把这条恒等式写死进契约。
    """
    kind = figure["kind"]
    if kind == "compare":
        tokens = []
        for side in ("left", "right"):
            tokens.append(figure[side]["title"])
            tokens.extend(figure[side]["items"])
        return tokens
    if kind == "matrix":
        return [*figure["columns"], *(cell for row in figure["rows"] for cell in row if cell)]
    if kind == "chart":
        return [item["label"] for item in figure["series"]]
    return list(figure["items"])


# =============================================================================
# 1. 六个渲染器都能画，且都画成可编辑形状
# =============================================================================

@pytest.mark.parametrize("kind", FIGURE_KINDS)
def test_every_figure_kind_renders_editable_shapes(kind):
    prs = Presentation(PPTExporter().export_to_bytes([_figure_node(kind)], "图示验收"))
    shapes = [shape for slide in prs.slides for shape in slide.shapes]

    assert shapes, f"{kind} 图示页一个形状都没有"
    assert all(s.shape_type != 13 for s in shapes), f"{kind} 导出了图片，教师改不动了"


@pytest.mark.parametrize("kind", FIGURE_KINDS)
def test_figure_and_conclusion_reach_both_ends(kind):
    """预览与 .pptx 必须显示同一份图示与同一句结论。"""
    node = _figure_node(kind)
    html = PreviewRenderer().render([node], "图示验收")
    ppt_text = _ppt_text([node])

    for token in _figure_tokens(node["figure"]):
        assert token in html, f"预览缺少图示标签：{token}"
        assert token in ppt_text, f"PPT 缺少图示标签：{token}"
    assert CONCLUSION in html
    assert CONCLUSION in ppt_text


# =============================================================================
# 2. 退化输入不能让整页导出崩掉
# =============================================================================

def test_two_node_cycle_degrades_to_a_flow():
    """两点"循环"在几何上就是一条来回的直线，硬摆圆环只会挤成一团。"""
    node = _figure_node("cycle")
    node["figure"] = {"kind": "cycle", "items": ["甲", "乙"]}

    prs = Presentation(PPTExporter().export_to_bytes([node], "图示验收"))

    assert [s.text for slide in prs.slides for s in slide.shapes if hasattr(s, "text")]
    assert len(prs.slides) == 1


def test_matrix_with_ragged_rows_still_renders():
    """模型经常给不齐列数；page_spec 已补齐，导出侧不能再自己数一遍列。"""
    node = _figure_node("matrix")
    node["figure"] = {"kind": "matrix", "columns": ["A", "B", "C"],
                      "rows": [["甲"], ["乙", "丙", "丁", "戊"]]}

    ppt_text = _ppt_text([node])

    assert "甲" in ppt_text and "丙" in ppt_text


def test_unknown_figure_kind_is_skipped_not_fatal():
    """未登记的 kind 由 page_spec 归一成"没有图示"，导出侧照常出正文。"""
    node = _figure_node("flow")
    node["figure"] = {"kind": "三维饼图", "items": ["甲", "乙"]}

    ppt_text = _ppt_text([node])

    assert "图示与结论必须两端同源。" in ppt_text
    assert CONCLUSION in ppt_text


# =============================================================================
# 3. 正文与要点并存时，要点不能消失
# =============================================================================

def test_key_points_survive_alongside_body_text():
    """n3 的形状：content 与 key_points 同时有值 → bullets 是 ROLE_SECONDARY。

    导出侧若按"只读 PRIMARY"实现，这两个要点会在 .pptx 里整块消失，而预览里
    还在，error 列表也是空的 —— 教师只会觉得导出的课件少了东西。
    """
    node = {
        "node_id": "n3", "type": "content", "node_type": "content",
        "title": "复杂度基础", "content": "复杂度描述算法规模增长趋势。",
        "key_points": ["时间复杂度", "空间复杂度"],
    }
    ppt_text = _ppt_text([node])

    assert "复杂度描述算法规模增长趋势。" in ppt_text
    assert "时间复杂度" in ppt_text
    assert "空间复杂度" in ppt_text


def test_point_details_reach_the_slide():
    """要点展开是本轮的主角 —— 它必须以「要点：展开」的形式真的上幻灯片。"""
    node = {
        "node_id": "n3", "type": "content", "node_type": "content",
        "title": "冒泡排序", "content": "每一轮把最大值送到末尾。",
        "point_details": [{"point": "相邻比较", "detail": "每次只比一对。"}],
    }
    ppt_text = _ppt_text([node])

    assert "相邻比较" in ppt_text
    assert "每次只比一对。" in ppt_text


# =============================================================================
# 4. 数值单位：不再替模型假设成百分比
# =============================================================================

def test_chart_values_are_not_forced_to_percent():
    """改动前 _render_bars 写死 "%"（当时数值来自百分比正则启发式，成立）。

    现在数值由规划器给，可以是人数、秒数、倍数 —— 再写死 "%" 在多数页面上
    就是错的。这条锁住那个刻意的观感变化。

    注意 fixture 必须给 ≥2 个数据点：单点图会被 page_spec.normalize_figure 判为
    "谈不上关系"而整块丢掉（见 test_single_series_chart_is_not_a_chart），那时页面上
    一个形状都没有，本测试就会以一个和单位毫无关系的理由失败。
    """
    node = _figure_node("chart")
    node["figure"] = {"kind": "chart", "series": [
        {"label": "冒泡排序", "value": 45}, {"label": "快速排序", "value": 12},
    ]}

    ppt_text = _ppt_text([node])

    assert "45" in ppt_text
    assert "12" in ppt_text
    assert "45%" not in ppt_text


def test_single_series_chart_is_not_a_chart():
    """一个数据点画不出"数据关系"，只有轨道和一根占满的柱子。

    这条门槛在 page_spec 里，不在导出侧 —— 写在这儿是为了让"图表没画出来"这类
    现象能一眼分清是归一器拒了，还是渲染器漏了。
    """
    node = _figure_node("chart")
    node["figure"] = {"kind": "chart", "series": [{"label": "冒泡排序", "value": 45}]}

    ppt_text = _ppt_text([node])

    assert "45" not in ppt_text
    # 正文与结论照常出：丢的只是图，不是整页
    assert "图示与结论必须两端同源。" in ppt_text
    assert CONCLUSION in ppt_text
