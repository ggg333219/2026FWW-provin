"""预览 HTML 与 PPTX 导出的同源语义一致性回归测试。"""

from pptx import Presentation

from ai_engine.gen.export_ppt import PPTExporter, resolve_node_type
from ai_engine.gen.presentation import presentation_from_nodes, presentation_to_theme
from ai_engine.gen.preview import PreviewRenderer


def _six_node_snapshot() -> list[dict]:
    return [
        {
            "node_id": "n1", "type": "cover", "node_type": "cover",
            "title": "排序算法课堂", "content": "六节点同源验收",
        },
        {
            "node_id": "n2", "type": "section", "node_type": "section",
            "title": "学习目录", "content": "从概念到实践",
            "key_points": ["理解复杂度", "比较典型算法"],
        },
        {
            "node_id": "n3", "type": "content", "node_type": "content",
            "title": "复杂度基础", "content": "复杂度描述算法规模增长趋势。",
            "key_points": ["时间复杂度", "空间复杂度"],
            "core_points": ["时间复杂度", "空间复杂度"],
            "segments": [
                {"segment_type": "opening", "text": "从查找成绩的耗时问题导入。", "estimated_seconds": 45},
                {"segment_type": "explanation", "text": "讲解大O表示法。", "estimated_seconds": 90},
            ],
        },
        {
            "node_id": "n4", "type": "two_column", "node_type": "two_column",
            "title": "算法对比", "content": "比较不同排序算法。",
            "key_points": ["冒泡排序稳定", "快速排序平均更快"],
            "core_points": ["冒泡排序稳定", "快速排序平均更快"],
            "segments": [
                {"segment_type": "example", "text": "对比O(n²)与O(n log n)。", "estimated_seconds": 60},
            ],
        },
        {
            "node_id": "n5", "type": "quiz", "node_type": "quiz",
            "title": "课堂检测", "questions": ["哪种算法平均复杂度为O(n log n)？"],
        },
        {
            "node_id": "n6", "type": "summary", "node_type": "summary",
            "title": "课程小结", "key_points": ["选择算法要结合数据规模", "关注时间与空间权衡"],
            "core_points": ["选择算法要结合数据规模", "关注时间与空间权衡"],
        },
    ]


def _ppt_text(nodes: list[dict]) -> tuple[Presentation, str]:
    prs = Presentation(PPTExporter().export_to_bytes(nodes, "六节点同源验收"))
    text = "\n".join(
        shape.text
        for slide in prs.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )
    return prs, text


def test_conflicting_type_resolution_matches_preview():
    node = {
        "node_id": "conflict", "title": "冲突兼容样本",
        "type": "quiz", "node_type": "cover", "questions": ["预览以哪个字段为准？"],
    }
    preview_type = PreviewRenderer()._resolve_type(node)
    assert preview_type == "quiz"
    assert resolve_node_type(node) == preview_type


def test_six_node_snapshot_keeps_authored_semantics_in_preview_and_ppt():
    nodes = _six_node_snapshot()
    preview_html = PreviewRenderer().render(nodes, "六节点同源验收")
    prs, ppt_text = _ppt_text(nodes)

    # 两端都必须出现的作者语义
    authored_tokens = [
        "排序算法课堂", "六节点同源验收", "学习目录", "从概念到实践",
        "复杂度基础", "复杂度描述算法规模增长趋势。", "时间复杂度", "空间复杂度",
        "算法对比", "比较不同排序算法。", "冒泡排序稳定", "快速排序平均更快",
        "课堂检测", "哪种算法平均复杂度为O(n log n)？",
        "课程小结", "选择算法要结合数据规模", "关注时间与空间权衡",
    ]
    for token in authored_tokens:
        assert token in preview_html, f"preview missing: {token}"
        assert token in ppt_text, f"ppt missing: {token}"

    # 讲稿分段只在预览出现：导出侧按需求不展示讲稿内容
    # （见 export_ppt.PPTExporter._segment_items）。这里双向断言，
    # 既防止有人把分段塞回 PPT，也防止预览丢掉讲稿。
    script_tokens = [
        "从查找成绩的耗时问题导入。", "讲解大O表示法。", "对比O(n²)与O(n log n)。",
    ]
    for token in script_tokens:
        assert token in preview_html, f"preview 丢失讲稿分段: {token}"
        assert token not in ppt_text, f"PPT 不应包含讲稿分段: {token}"

    # M3 默认结构：作者语义页数不变，仅末尾追加"感谢聆听"致谢页（6 + 1 = 7）
    assert len(prs.slides) == 7
    assert all(shape.shape_type != 13 for slide in prs.slides for shape in slide.shapes)
    assert resolve_node_type(nodes[3]) == "two_column"


def test_preview_applies_presentation_style_from_nodes():
    nodes = [{
        "node_id": "styled",
        "type": "content",
        "title": "样式预览",
        "content": "样式应反映到实时预览。",
        "metadata": {"presentation": {
            "palette": "klein-blue",
            "layout": "text",
            "font_family": "serif",
            "density": "minimal",
        }},
    }]

    html = PreviewRenderer().render(nodes, "样式预览")

    assert 'data-palette="klein-blue"' in html
    assert "font-family: 'Songti SC'" in html
    assert "--preview-slide-padding: 56px 64px 60px" in html


def _styled_six_node_snapshot(palette: str = "klein-blue",
                             font_family: str = "serif") -> list[dict]:
    """带样式层的六节点快照；样式挂在首个节点上（与 Go 侧 presentationFromNodes 一致）。"""
    nodes = _six_node_snapshot()
    nodes[0] = {**nodes[0], "metadata": {"presentation": {
        "palette": palette, "font_family": font_family,
    }}}
    return nodes


def _hex(color) -> str:
    """RGBColor → "RRGGBB"。不依赖 python-pptx 的 __str__ 实现，直接按字节格式化。"""
    return "%02X%02X%02X" % (color[0], color[1], color[2])


def _solid_fills(prs: Presentation) -> set[str]:
    """收集所有纯色填充；渐变/无填充的形状自动跳过。"""
    colors = set()
    for slide in prs.slides:
        for shape in slide.shapes:
            try:
                colors.add(_hex(shape.fill.fore_color.rgb))
            except Exception:
                continue
    return colors


def _run_fonts(prs: Presentation) -> set[str]:
    fonts = set()
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.name:
                        fonts.add(run.font.name)
    return fonts


def test_export_follows_presentation_style():
    """样式中选定的配色与字体必须落到导出的 PPTX 上。"""
    prs = Presentation(PPTExporter().export_to_bytes(_styled_six_node_snapshot(), "样式导出"))

    assert "315FD4" in _solid_fills(prs), "Klein Blue 未落到形状填充"
    assert "SimSun" in _run_fonts(prs), "serif 字体族未落到文本"
    # M3 默认结构：追加致谢页（6 + 1 = 7），样式断言不受影响
    assert len(prs.slides) == 7
    assert all(shape.shape_type != 13 for slide in prs.slides for shape in slide.shapes)


def test_export_legacy_theme_matches_preview():
    """没传样式层的历史课件，配色必须与预览同一套（classic 调色板）。

    改动前这条断言的是 catalog.default 的 2D5F58 —— 预览从来不用 catalog，
    于是"没选样式"的课件预览是亮青绿、导出是暗墨绿，八个令牌没一个对得上。
    预览是样式的单一事实来源，导出跟着它走。
    """
    prs = Presentation(PPTExporter().export_to_bytes(_six_node_snapshot(), "旧课件"))
    fills = _solid_fills(prs)

    assert "168F82" in fills, "历史课件未使用预览的 classic 主色"
    assert "315FD4" not in fills, "历史课件被误加了别的调色板"


def test_presentation_to_theme_falls_back_to_preview_palette():
    """取值顺序：显式 palette → 模板自带 palette → 默认调色板 classic。"""
    theme, styled = presentation_to_theme(*presentation_from_nodes([{}]), "default")
    assert styled is True and theme.primary == "168F82"

    styled_nodes = [{"metadata": {"presentation": {"palette": "klein-blue"}}}]
    theme, styled = presentation_to_theme(*presentation_from_nodes(styled_nodes), "default")
    assert styled is True and theme.primary == "315FD4"

    # 选了自带样式的模板但前端没给 palette → 用该模板的默认配色
    theme, styled = presentation_to_theme(*presentation_from_nodes([{}]), "deck-swiss-edu")
    assert styled is True and theme.primary == "315FD4"

    # 自带配色的 catalog 主题（academic）仍保留自己的令牌，不让位给 classic
    theme, styled = presentation_to_theme(*presentation_from_nodes([{}]), "academic")
    assert styled is True and theme.primary == "2D5F58"
