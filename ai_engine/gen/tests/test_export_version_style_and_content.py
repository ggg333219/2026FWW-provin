"""Version export must honor saved style and never invent interaction content."""

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from ai_engine.gen.deck_structure import expand_deck_pages
from ai_engine.gen.export_ppt import ExportPPTConfig, PPTExporter
from ai_engine.gen.presentation import apply_export_presentation, content_geometry


def _exporter():
    return PPTExporter(ExportPPTConfig(
        include_auto_cover=False,
        include_auto_toc=False,
        include_auto_summary=False,
        include_closing_slide=False,
    ))


def _content(style):
    return {
        "node_id": "content-1", "type": "content", "title": "排序算法",
        "content": "快速排序采用分治思想", "key_points": ["划分数组"],
        "metadata": {"presentation": style},
    }


def test_separate_version_style_is_attached_without_changing_snapshot():
    original = [{"node_id": "n1", "type": "content", "title": "排序算法",
                 "metadata": {"interactive_asset": {"items": []}}}]
    styled = apply_export_presentation(original, {"font_family": "serif", "layout": "text"})
    assert styled[0]["metadata"]["presentation"]["font_family"] == "serif"
    assert styled[0]["metadata"]["interactive_asset"] == {"items": []}
    assert "presentation" not in original[0]["metadata"]


def test_east_asian_font_follows_saved_font_family():
    for family, expected in (("serif", "SimSun"), ("source-han", "Source Han Sans SC")):
        prs = Presentation(_exporter().export_to_bytes([_content({"font_family": family})]))
        east_asian = {
            ea.get("typeface")
            for slide in prs.slides for shape in slide.shapes
            if getattr(shape, "has_text_frame", False)
            for paragraph in shape.text_frame.paragraphs for run in paragraph.runs
            for ea in [run.font._rPr.find(qn("a:ea"))] if ea is not None
        }
        assert expected in east_asian
        assert "微软雅黑" not in east_asian


def test_layout_mode_changes_exported_content_panels():
    positions = []
    for mode in ("balanced", "visual", "text"):
        prs = Presentation(_exporter().export_to_bytes([_content({"layout": mode})]))
        slide = prs.slides[0]
        left_x, left_width = content_geometry("comfortable", mode).left
        assert any(abs(shape.left - Inches(left_x)) < Inches(.01)
                   and abs(shape.width - Inches(left_width)) < Inches(.01)
                   for shape in slide.shapes)
        positions.append(left_width)
    assert len(set(positions)) == 3


def test_quiz_uses_saved_asset_questions_instead_of_placeholder():
    node = {
        "node_id": "quiz-1", "type": "quiz", "title": "课堂检测",
        "metadata": {"interactive_asset": {"items": [
            {"type": "quiz", "questions": [{
                "question": "快速排序的平均时间复杂度是多少？",
                "options": ["O(n²)", "O(n log n)"], "answer": "B",
            }]},
        ]}},
    }
    prs = Presentation(_exporter().export_to_bytes([node]))
    text = "\n".join(shape.text for slide in prs.slides for shape in slide.shapes
                     if getattr(shape, "has_text_frame", False))
    assert "快速排序的平均时间复杂度是多少？" in text
    assert "请结合本节内容提出一个关键问题" not in text


def test_missing_quiz_and_plan_only_flashcard_fail_instead_of_mocking():
    with pytest.raises(ValueError, match="没有已保存的真实题目"):
        _exporter().export_to_bytes([{"node_id": "q", "type": "quiz", "title": "课堂检测"}])
    with pytest.raises(ValueError, match="只有规划、没有已保存的卡片内容"):
        _exporter().export_to_bytes([{
            "node_id": "f", "type": "interactive", "title": "知识翻卡",
            "interactive_plan": {"kind": "flashcard"}, "key_points": ["分治思想"],
        }])


def test_auto_summary_uses_only_teaching_points_and_omits_empty_placeholder():
    pages = expand_deck_pages([
        {"node_id": "c", "type": "content", "title": "快速排序", "key_points": ["分治思想"]},
        {"node_id": "q", "type": "quiz", "title": "课堂检测", "key_points": ["测试占位"]},
    ], include_cover=False, include_toc=False)
    summary = pages[-1]
    assert summary["node_id"] == "__auto_summary__"
    assert summary["key_points"] == ["分治思想"]
    assert "回顾本次课程的核心要点" not in str(summary)

    no_summary = expand_deck_pages([
        {"node_id": "q", "type": "quiz", "title": "课堂检测"},
    ], include_cover=False, include_toc=False)
    assert all(node.get("node_id") != "__auto_summary__" for node in no_summary)


def test_summary_density_changes_cards_without_fixed_filler():
    widths = []
    for density in ("compact", "comfortable", "minimal"):
        node = {
            "node_id": "summary-1", "type": "summary", "title": "本节小结",
            "key_points": ["快速排序采用分治思想"],
            "metadata": {"presentation": {"density": density}},
        }
        prs = Presentation(_exporter().export_to_bytes([node]))
        slide = prs.slides[0]
        text = "\n".join(shape.text for shape in slide.shapes
                         if getattr(shape, "has_text_frame", False))
        assert "快速排序采用分治思想" in text
        assert "感谢聆听 · 本节内容到此结束" not in text
        card = next(shape for shape in slide.shapes
                    if abs(shape.width - Inches(content_geometry(density).wide_right[1])) < Inches(.01)
                    and shape.height > Inches(.6))
        widths.append(card.width)
    assert len(set(widths)) == 3
