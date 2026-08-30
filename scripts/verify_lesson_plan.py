# -*- coding: utf-8 -*-
"""Word 标准教案链路离线验证脚本 (verify_lesson_plan.py)

不调用真实 LLM API（当日额度已耗尽，全部走 mock / monkeypatch / fallback）：
  T1  AI_GEN_MODE=mock            → 直接构造样例教案，结构完整
  T2  LLM 正常返回（MockLLMClient）→ 走 llm 路径，内容与预设一致
  T3  LLM chat_with_schema 抛异常（monkeypatch，模拟 429/超时）→ 确定性 fallback
  T4  LLM 返回 error 标记 / 空教案  → 确定性 fallback
  T5  Pydantic TeachingNode + IntentSummary 输入 → fallback 正常
  T6  DOCX 版式（mock 教案）：含 教学目标/教学方法/教学过程/课后作业 章节、
      教学过程表 >=4 列（实际 5 列）、无空章节、中文字体微软雅黑(eastAsia)
  T7  DOCX 版式（fallback 教案）：同上（LLM 失败路径导出也完整）
  T8  ExportFormat.LESSON_PLAN 枚举注册，值 "lesson_plan"
  T9  main._run_export 集成：fmt=lesson_plan → 任务 done、doc_path 产出真实 docx、
      记录 lesson_plan_source

运行：python scripts/verify_lesson_plan.py   （退出码 0 = 全部 PASS）
"""

import asyncio
import io
import os
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

os.environ.setdefault("PYTHONIOENCODING", "utf-8")

RESULTS = []


def check(name, condition, detail=""):
    status = "PASS" if condition else "FAIL"
    RESULTS.append((name, status, detail))
    print(f"[{status}] {name}" + (f" — {detail}" if detail else ""))


def make_sample_nodes():
    """标准样例节点（dict 形式，对齐版本快照 nodes_snapshot）。"""
    return [
        {"node_id": "n1", "type": "title", "title": "排序算法",
         "content": "课程封面：排序算法入门"},
        {"node_id": "n2", "type": "content", "title": "冒泡排序原理",
         "content": "相邻元素两两比较，每一轮将当前最大值沉底。",
         "core_points": ["两两比较", "每轮确定一个最大值"],
         "estimated_duration_minutes": 15},
        {"node_id": "n3", "type": "content", "title": "快速排序思想",
         "content": "选定基准元素，分而治之，递归排序两侧子序列。",
         "core_points": ["基准选取", "分区操作"],
         "estimated_duration_minutes": 15},
        {"node_id": "n4", "type": "quiz", "title": "随堂练习",
         "questions": ["对序列 [5,3,8] 手工模拟一轮冒泡排序。",
                       "冒泡排序在最好情况下的时间复杂度是什么？"]},
        {"node_id": "n5", "type": "summary", "title": "课堂小结",
         "content": "回顾冒泡排序与快速排序的核心思想与适用场景。"},
    ]


SAMPLE_INTENT = {
    "topic": "排序算法",
    "duration_minutes": 45,
    "style": "讲练结合",
    "target_student": "大一计算机专业",
    "knowledge_points": ["冒泡排序", "快速排序"],
    "with_interactive": True,
    "key_difficulties": ["递归分区边界的理解"],
}


def plan_structure_ok(plan, source=None):
    """教案结构完整性断言（章节齐全且无空章节）。"""
    problems = []
    if source is not None and plan.data_source != source:
        problems.append(f"data_source={plan.data_source} 期望 {source}")
    if not plan.course_name:
        problems.append("course_name 为空")
    if plan.duration_minutes <= 0:
        problems.append("duration_minutes 非法")
    dims = {o.get("dimension") for o in plan.teaching_objectives}
    for dim in ("知识目标", "能力目标", "素养目标"):
        if dim not in dims:
            problems.append(f"缺少 {dim}")
    for o in plan.teaching_objectives:
        if not str(o.get("content") or "").strip():
            problems.append("存在空目标")
    kd = plan.key_difficulties or {}
    if not (kd.get("key_points") or kd.get("difficult_points")):
        problems.append("重点难点为空")
    if not plan.teaching_methods:
        problems.append("teaching_methods 为空")
    phases = plan.teaching_process
    if len(phases) < 3:
        problems.append(f"教学过程阶段不足: {len(phases)}")
    phase_names = {p.phase for p in phases}
    for name in ("导入", "新授", "小结"):
        if name not in phase_names:
            problems.append(f"教学过程缺少 {name} 阶段")
    for p in phases:
        if not str(p.teacher_activity or "").strip():
            problems.append(f"阶段 {p.phase} 教师活动为空")
        if not str(p.student_activity or "").strip():
            problems.append(f"阶段 {p.phase} 学生活动为空")
        if not p.duration_minutes or p.duration_minutes <= 0:
            problems.append(f"阶段 {p.phase} 时长非法")
    if not plan.classroom_activities:
        problems.append("classroom_activities 为空")
    if not plan.homework:
        problems.append("homework 为空")
    return problems


def read_docx_full_text(buf: io.BytesIO) -> tuple:
    from docx import Document
    doc = Document(buf)
    paras = [p.text for p in doc.paragraphs]
    table_cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]
    return doc, "\n".join(paras + table_cells)


def docx_layout_ok(buf: io.BytesIO) -> list:
    """DOCX 版式断言，返回问题列表。"""
    problems = []
    doc, full = read_docx_full_text(buf)
    for section in ("教学目标", "教学方法", "教学过程", "课后作业"):
        if section not in full:
            problems.append(f"缺少章节: {section}")
    # 基本信息表
    if not doc.tables or "课程名称" not in doc.tables[0].rows[0].cells[0].text:
        problems.append("基本信息表缺失或表头不符")
    # 教学过程表：找表头为"阶段"的表，列数 >= 4
    proc_tables = [t for t in doc.tables if t.rows and t.rows[0].cells[0].text.strip() == "阶段"]
    if not proc_tables:
        problems.append("教学过程表缺失")
    else:
        cols = len(proc_tables[0].columns)
        if cols < 4:
            problems.append(f"教学过程表仅 {cols} 列 (<4)")
        headers = [proc_tables[0].rows[0].cells[i].text.strip() for i in range(cols)]
        for h in ("教师活动", "学生活动", "设计意图"):
            if h not in headers:
                problems.append(f"教学过程表缺少列: {h}")
        # 无空章节/空单元格：过程表数据行不得整行为空
        for r in proc_tables[0].rows[1:]:
            if not any(c.text.strip() for c in r.cells):
                problems.append("教学过程表存在空行")
    # 中文字体微软雅黑（eastAsia）
    from docx.oxml.ns import qn
    title_runs = doc.paragraphs[0].runs
    if not title_runs:
        problems.append("标题段落无 run")
    else:
        rfonts = title_runs[0]._element.rPr.rFonts
        ea = rfonts.get(qn("w:eastAsia")) if rfonts is not None else None
        if ea != "微软雅黑":
            problems.append(f"标题 eastAsia 字体为 {ea!r}，期望 微软雅黑")
    return problems


class FailingLLMClient:
    """模拟 LLM 彻底不可用（429 / 超时 / 网络异常）。"""

    def chat_with_schema(self, *args, **kwargs):
        raise RuntimeError("mock: LLM 429 quota exhausted")

    def chat_json(self, *args, **kwargs):
        raise RuntimeError("mock: LLM 429 quota exhausted")


class ErrorPayloadLLMClient:
    """模拟 LLM 返回带 error 标记 / 结构缺失的 JSON。"""

    def __init__(self, payload):
        self.payload = payload

    def chat_with_schema(self, *args, **kwargs):
        return self.payload


def _set_mode(settings, value):
    """Settings 为 frozen dataclass，测试内通过 object.__setattr__ 切换 GEN_MODE。"""
    object.__setattr__(settings, "GEN_MODE", value)


def main():
    from ai_engine.core.config import settings
    from ai_engine.core.schema import ExportFormat
    from ai_engine.gen.lesson_plan import LessonPlanGenerator
    from ai_engine.gen.export_doc import LessonPlanDOCExporter

    old_mode = settings.GEN_MODE

    nodes = make_sample_nodes()
    intent = dict(SAMPLE_INTENT)

    # ---- T1: mock 模式直接构造 ----
    _set_mode(settings, "mock")
    plan = LessonPlanGenerator().generate(nodes, intent)
    problems = plan_structure_ok(plan, source="mock")
    check("T1 mock 模式构造完整教案", not problems, "; ".join(problems))

    # ---- T2: LLM 正常返回 → llm 路径 ----
    from ai_engine.gen.tests.mock_llm import MockLLMClient
    _set_mode(settings, "llm")
    preset = {
        "course_name": "排序算法",
        "subject": "计算机科学",
        "grade_level": "大一计算机专业",
        "duration_minutes": 45,
        "teaching_objectives": [
            {"dimension": "知识目标", "content": "掌握冒泡与快排的核心思想。"},
            {"dimension": "能力目标", "content": "能手工模拟排序过程。"},
            {"dimension": "素养目标", "content": "培养算法思维。"},
        ],
        "key_difficulties": {"key_points": ["两种排序的对比"], "difficult_points": ["递归分区"]},
        "teaching_methods": ["讲授法", "任务驱动法"],
        "teaching_process": [
            {"phase": "导入", "duration_minutes": 5, "teacher_activity": "情境导入",
             "student_activity": "思考回答", "design_intent": "激发兴趣"},
            {"phase": "新授", "duration_minutes": 25, "teacher_activity": "讲解原理",
             "student_activity": "记录要点", "design_intent": "建构知识"},
            {"phase": "练习", "duration_minutes": 10, "teacher_activity": "布置练习",
             "student_activity": "动手模拟", "design_intent": "巩固应用"},
            {"phase": "小结", "duration_minutes": 5, "teacher_activity": "梳理脉络",
             "student_activity": "回顾总结", "design_intent": "强化记忆"},
        ],
        "classroom_activities": [{"name": "排序模拟赛", "description": "分组竞赛模拟排序"}],
        "homework": ["完成教材第3章习题 1-3 题"],
        "board_design": "课题：排序算法",
    }
    mock_client = MockLLMClient()
    mock_client.set_json_response(preset)
    plan_llm = LessonPlanGenerator(llm_client=mock_client).generate(nodes, intent)
    ok = (plan_llm.data_source == "llm"
          and plan_llm.course_name == "排序算法"
          and plan_llm.teaching_objectives[0]["content"] == "掌握冒泡与快排的核心思想。"
          and plan_llm.teaching_process[1].duration_minutes == 25)
    check("T2 LLM 正常返回走 llm 路径且内容一致", ok,
          f"source={plan_llm.data_source}, title={plan_llm.course_name}")

    # ---- T3: LLM 抛异常 → fallback 完整教案 ----
    plan_fb = LessonPlanGenerator(llm_client=FailingLLMClient()).generate(nodes, intent)
    problems = plan_structure_ok(plan_fb, source="fallback")
    check("T3 LLM 异常走确定性 fallback 且结构完整", not problems, "; ".join(problems))
    check("T3 fallback 课时取自意图", plan_fb.duration_minutes == 45,
          str(plan_fb.duration_minutes))
    check("T3 fallback 知识点进入目标",
          any("冒泡排序" in o["content"] for o in plan_fb.teaching_objectives),
          plan_fb.teaching_objectives[0]["content"][:40])
    check("T3 fallback 互动问题进入作业",
          any("冒泡" in h or "复杂度" in h for h in plan_fb.homework),
          str(plan_fb.homework[:1]))

    # ---- T4: LLM 返回 error 标记 / 空结构 → fallback ----
    plan_e1 = LessonPlanGenerator(
        llm_client=ErrorPayloadLLMClient({"error": "json_parse_failed"})
    ).generate(nodes, intent)
    problems = plan_structure_ok(plan_e1, source="fallback")
    check("T4a LLM error 标记 → fallback", not problems, "; ".join(problems))
    plan_e2 = LessonPlanGenerator(
        llm_client=ErrorPayloadLLMClient({"course_name": "x", "teaching_process": []})
    ).generate(nodes, intent)
    problems = plan_structure_ok(plan_e2, source="fallback")
    check("T4b LLM 空教案结构 → fallback", not problems, "; ".join(problems))

    # ---- T5: Pydantic 输入（TeachingNode + IntentSummary）----
    from ai_engine.core.schema import TeachingNode, IntentSummary
    pyd_nodes = [TeachingNode(node_id=f"p{i}", title=n["title"], content=n.get("content", ""),
                              core_points=n.get("core_points", []),
                              interactive_questions=n.get("questions", []),
                              estimated_duration_minutes=10)
                 for i, n in enumerate(make_sample_nodes())]
    plan_pyd = LessonPlanGenerator().generate(pyd_nodes, IntentSummary(**SAMPLE_INTENT))
    problems = plan_structure_ok(plan_pyd, source="fallback")
    check("T5 Pydantic 输入 fallback 正常", not problems, "; ".join(problems))

    _set_mode(settings, old_mode)

    # ---- T6: mock 教案 DOCX 版式 ----
    buf_mock = LessonPlanDOCExporter().export_to_bytes(plan)
    problems = docx_layout_ok(buf_mock)
    check("T6 mock 教案 DOCX 版式合规", not problems, "; ".join(problems))

    # ---- T7: fallback 教案 DOCX 版式（LLM 失败路径导出也完整）----
    buf_fb = LessonPlanDOCExporter().export_to_bytes(plan_fb)
    problems = docx_layout_ok(buf_fb)
    check("T7 fallback 教案 DOCX 版式合规", not problems, "; ".join(problems))

    # ---- T8: ExportFormat 枚举注册 ----
    check("T8 ExportFormat.LESSON_PLAN 已注册",
          hasattr(ExportFormat, "LESSON_PLAN") and ExportFormat.LESSON_PLAN.value == "lesson_plan",
          getattr(getattr(ExportFormat, "LESSON_PLAN", None), "value", "missing"))

    # ---- T9: main._run_export 集成（fmt=lesson_plan）----
    from ai_engine import main as app_main
    from ai_engine.core.schema import ExportJob, ExportStatus

    _set_mode(settings, "mock")  # 集成用例不触真实 LLM

    async def run_case():
        job_id = "exp_lp_t9"
        job = ExportJob(job_id=job_id, generation_id="gen_lp", format=ExportFormat.LESSON_PLAN,
                        status=ExportStatus.PENDING)
        app_main._export_tasks[job_id] = job
        await app_main._run_export(job_id, nodes, ExportFormat.LESSON_PLAN, intent=intent)
        return app_main._export_tasks.pop(job_id)

    job = asyncio.run(run_case())
    meta = job.metadata or {}
    check("T9 导出任务 done", job.status == ExportStatus.DONE, str(job.status))
    doc_path = meta.get("doc_path", "")
    check("T9 产出 docx 到 doc_path", doc_path.endswith(".docx") and Path(doc_path).is_file(),
          doc_path)
    check("T9 记录教案来源", meta.get("lesson_plan_source") == "mock",
          str(meta.get("lesson_plan_source")))
    if doc_path and Path(doc_path).is_file():
        doc, full = read_docx_full_text(io.BytesIO(Path(doc_path).read_bytes()))
        ok = all(k in full for k in ("教学目标", "教学过程", "教学方法", "课后作业"))
        check("T9 导出 docx 内容可读且章节齐全", ok)
    else:
        check("T9 导出 docx 内容可读且章节齐全", False, "docx 未产出")

    _set_mode(settings, old_mode)

    # ---- 汇总 ----
    total = len(RESULTS)
    passed = sum(1 for _, s, _ in RESULTS if s == "PASS")
    print(f"\n===== verify_lesson_plan: {passed}/{total} PASS =====")
    failed = [(n, d) for n, s, d in RESULTS if s == "FAIL"]
    for name, detail in failed:
        print(f"  FAIL: {name} — {detail}")
    sys.exit(0 if not failed else 1)


if __name__ == "__main__":
    main()
