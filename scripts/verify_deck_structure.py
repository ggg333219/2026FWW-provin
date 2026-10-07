"""M3 验证：默认 PPT 输出必须包含 封面 / 目录（内容节点>5 时）/ 总结页。

验证口径（预览与导出同源）：
  1. 8 个内容节点导出 PPTX：第1页封面（含课名）、第2页目录、末2页 = 本节小结 + 感谢聆听
  2. 4 个内容节点（≤5）：不注入目录页
  3. 首节点已是封面：不重复注入封面，目录页紧随其后
  4. 空节点：导出永不空白（至少有封面 + 总结）
  5. 预览 HTML：首个 slide 为 cover-slide，含目录页与总结页

用法：conda run -n fuww_ai python scripts/verify_deck_structure.py
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from ai_engine.gen.export_ppt import ExportPPTConfig, PPTExporter  # noqa: E402
from ai_engine.gen.preview import PreviewRenderer  # noqa: E402
from pptx import Presentation  # noqa: E402

PASS, FAIL = "✅", "❌"
results: list[tuple[bool, str]] = []


def check(ok: bool, message: str) -> None:
    results.append((bool(ok), message))
    print(f"{PASS if ok else FAIL} {message}")


def slide_texts(prs: Presentation) -> list[str]:
    texts = []
    for slide in prs.slides:
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
        texts.append(" ".join(parts))
    return texts


def make_nodes(titles: list[str], first_cover: bool = False) -> list[dict]:
    nodes = []
    if first_cover:
        nodes.append({"node_id": "node_cover", "title": "自定义封面", "node_type": "cover",
                      "content": "自定义副标题", "key_points": []})
    for i, title in enumerate(titles, start=1):
        nodes.append({
            "node_id": f"node_{i:03d}", "title": title, "node_type": "content",
            "content": f"{title}的讲解内容。",
            "key_points": [f"{title}要点", "通用要点"],
        })
    return nodes


def export_nodes(nodes: list[dict], course_name: str) -> Presentation:
    exporter = PPTExporter(ExportPPTConfig(output_dir=str(Path(tempfile.gettempdir()) / "a12_m3")))
    path = exporter.export(nodes, course_name)["path"]
    return Presentation(path)


TITLES_8 = [f"第{i}节" for i in range(1, 9)]  # 8 个内容节点（含要点）
TITLES_4 = [f"第{i}节" for i in range(1, 5)]


def main() -> int:
    # ── 场景 1：8 内容节点 → 封面 + 目录2页(7条/页) + 8 内容 + 总结 + 感谢聆听 = 13 页 ──
    prs = export_nodes(make_nodes(TITLES_8), "排序算法")
    texts = slide_texts(prs)
    check(len(prs.slides) == 13, f"场景1 页数=13（实际 {len(prs.slides)}）：封面+目录2页+8内容+小结+致谢")
    check("排序算法" in texts[0], f"场景1 首页为封面且含课名（首页文本: {texts[0][:40]!r}）")
    check("目录" in texts[1] and "第1节" in texts[1], f"场景1 第2页为目录且含条目（{texts[1][:40]!r}）")
    check("本节小结" in texts[-2], f"场景1 倒数第2页为总结页（{texts[-2][:40]!r}）")
    check("感谢聆听" in texts[-1], "场景1 末页为致谢页")
    check("第8节" in texts[-3], "场景1 内容页全部保留（第8节在总结页之前）")

    # ── 场景 2：4 内容节点（≤5）→ 无目录页 ──
    prs2 = export_nodes(make_nodes(TITLES_4), "小课件")
    texts2 = slide_texts(prs2)
    check(len(prs2.slides) == 7, f"场景2 页数=7（实际 {len(prs2.slides)}）：封面+4内容+小结+致谢，无目录")
    check(not any("目录" in t and "第1节" in t for t in texts2), "场景2 未注入目录页")

    # ── 场景 3：首节点已是封面 → 不重复注入，目录紧随其后 ──
    prs3 = export_nodes(make_nodes(TITLES_8, first_cover=True), "自定义封面课件")
    texts3 = slide_texts(prs3)
    check(len(prs3.slides) == 13, f"场景3 页数=13（实际 {len(prs3.slides)}）：已有封面不再注入")
    check("自定义封面" in texts3[0], "场景3 首页保留用户自己的封面")
    check("目录" in texts3[1] and "第1节" in texts3[1], "场景3 目录页紧随用户封面")

    # ── 场景 4：空节点 → 永不空白 ──
    prs4 = export_nodes([], "空课件")
    texts4 = slide_texts(prs4)
    check(len(prs4.slides) >= 2, f"场景4 空节点导出仍≥2页（实际 {len(prs4.slides)}）")
    check("空课件" in texts4[0], "场景4 空节点也生成封面")

    # ── 场景 5：预览 HTML 与导出同源 ──
    html = PreviewRenderer().render(make_nodes(TITLES_8), "排序算法")
    first_slide_ok = "cover-slide" in html and html.index("cover-slide") < html.index("section-slide")
    check(first_slide_ok, "场景5 预览首个幻灯片为封面版式")
    check("目录" in html and "本节小结" in html, "场景5 预览含目录页与总结页（与导出同源）")
    slide_count = html.count('class="slide')
    check(slide_count >= 11, f"场景5 预览幻灯片数量≥11（实际 {slide_count}）")

    # ── 总结：真实节点末尾已有 summary 时不再重复注入 ──
    nodes_with_summary = make_nodes(TITLES_4)
    nodes_with_summary.append({
        "node_id": "node_summary", "title": "课程回顾", "node_type": "summary",
        "content": "", "key_points": ["要点一", "要点二"],
    })
    prs6 = export_nodes(nodes_with_summary, "带总结课件")
    texts6 = slide_texts(prs6)
    check(sum(1 for t in texts6 if "本节小结" in t) == 0, "场景6 末尾已有总结页时不重复注入'本节小结'")
    check("课程回顾" in texts6[-2], "场景6 保留用户自己的总结页")

    passed = sum(1 for ok, _ in results if ok)
    print(f"\n===== 验证结果：{passed}/{len(results)} 通过 =====")
    return 0 if passed == len(results) else 1


if __name__ == "__main__":
    raise SystemExit(main())
